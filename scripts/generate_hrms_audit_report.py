"""
Generate HRMS_Audit_Report.docx — Phases 1–3C compiled into stakeholder DOCX.
Title: HRMS Functionality, Role, Industry & Security Audit Report
Method: Static code-path tracing (UI → route → controller → model). Brutally honest.
"""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent.parent / "HRMS_Audit_Report.docx"

OK = "Fully Working"
PARTIAL = "Partially Working"
BROKEN = "Not Working"
MANUAL = "Needs Manual Testing"
STATUS_SYM = {
    OK: "✅",
    PARTIAL: "⚠️",
    BROKEN: "❌",
    MANUAL: "⚠️ Needs manual testing",
}
STATUS_FILL = {
    OK: "C6EFCE",
    PARTIAL: "FFEB9C",
    BROKEN: "FFC7CE",
    MANUAL: "FFEB9C",
}
GAP_FILL = {
    "Fully Addressed": "C6EFCE",
    "Partially Addressed": "FFEB9C",
    "Not Addressed": "FFC7CE",
}
SEC_FILL = {
    "✅": "C6EFCE",
    "⚠️": "FFEB9C",
    "❌": "FFC7CE",
}
YES_FILL = {"Yes": "C6EFCE", "Partial": "FFEB9C", "No": "FFC7CE"}


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def style_run(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, center=False, color=None, space_after=6):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_table(doc, headers, rows, status_col=None, status_map=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "1A365D")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    fill_map = status_map or STATUS_FILL
    for row_data in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = str(val) if val is not None else "—"
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7)
            if status_col is not None and i == status_col:
                raw = str(val)
                shaded = False
                for key, fill in fill_map.items():
                    if key in raw:
                        set_cell_shading(cells[i], fill)
                        shaded = True
                        break
                if not shaded:
                    if "✅" in raw or raw == "Yes":
                        set_cell_shading(cells[i], "C6EFCE")
                    elif "❌" in raw or raw == "No":
                        set_cell_shading(cells[i], "FFC7CE")
                    elif "⚠️" in raw or raw == "Partial":
                        set_cell_shading(cells[i], "FFEB9C")
    doc.add_paragraph()
    return table


def sym(status):
    return f"{STATUS_SYM.get(status, '')} {status}".strip()


# =============================================================================
# PHASE 1 — Inventory
# =============================================================================
INVENTORY = {
    "Authentication & Session": [
        ("Unified login", "Email/password login for admin/employee/org; role dashboard redirect", "Auth/Login.jsx → POST /api/v1/both/auth/login → unifiedAuthController", "Public"),
        ("Logout", "Clears JWT/refresh/CSRF cookies and Redux auth state", "Header → authSlice.logoutUser", "Any authenticated"),
        ("Organization signup", "Multi-step org + admin registration", "Auth/OrganizationRegister.jsx → POST /v1/organization/auth/register", "Public"),
        ("Organization login", "Org-owner login", "Auth/OrganizationLogin.jsx", "Public"),
        ("Admin register (dev)", "Admin account creation", "Auth/Register.jsx → POST /v1/admin/auth/register", "Gated in production"),
        ("Employee register (HR-created)", "Create employee credentials during hire", "AddEmployee → POST /v1/employee/auth/register", "admin, hr + employee.create"),
        ("Forgot / reset password", "Email token reset flow", "ForgotPasswordEmail / ResetPassword → /password/*", "Public (token)"),
        ("Session refresh", "Access ~1h + refresh cookie; axios 401→refresh", "axios.js + refresh-token routes", "Any authenticated"),
        ("Frontend route guards", "Role/permission-gated SPA routes", "ProtectedRoute.jsx + routes/index.jsx", "All roles"),
        ("API auth middleware", "JWT + requirePermission on many APIs", "middlewares/* + requirePermission", "Server-side"),
        ("First-login permissions", "Location/push permission prompt", "FirstLoginPermissionRequest.jsx", "Any authenticated"),
    ],
    "Dashboards": [
        ("Role-based dashboard redirect", "Picks dashboard from role/API", "RoleBasedDashboard.jsx → /dashboard", "Any authenticated"),
        ("Admin / HR / Manager / Employee / Finance / Tech / Executive dashboards", "Role-specific KPI hubs", "Admin|HR|Manager|Employee|Finance|Tech|Administration Dashboard*.jsx", "Per-route roles"),
        ("Organization dashboard", "Org-level overview", "OrganizationDashboard.jsx", "admin, executive"),
        ("Employee todos widget", "Personal todo list (localStorage only)", "Employee/Dashboard/Todo.jsx", "Any authenticated"),
    ],
    "Employee Management": [
        ("Employee list / search", "Search/filter workforce", "EmployeesPage.jsx → GET /admin/employee", "admin, hr, manager, executive"),
        ("Add employee", "Create employee + credentials + system role", "AddEmployee.jsx → employee auth register", "admin, hr + employee.create"),
        ("Employee details (multi-tab)", "Personal/job/docs/leave/payroll/performance/timeline", "EmployeeDetailsPage.jsx", "admin, hr, manager, executive"),
        ("Edit work info", "Update department/title/manager", "PUT /admin/employee/:id/work-info", "admin, hr, manager (scoped)"),
        ("Self profile view/edit", "Own profile fields + tax/financial", "ViewProfile / ProfileDetails / TaxDeclaration", "Any authenticated (self)"),
        ("Soft deactivate / terminate", "Status lifecycle without hard wipe", "PATCH update-status / terminate", "admin, hr"),
        ("Hard delete employee", "Permanently remove employee document", "DELETE .../delete-employee/:id", "admin"),
        ("Profile photo upload", "Upload avatar", "ProfilePhotoRoutes + multer", "Self / admin (upload auth'd)"),
        ("Bulk import / export", "CSV hire upload / Excel export", "— (CSV mislabeled export only)", "—"),
    ],
    "Attendance": [
        ("Clock in / out + breaks", "Punch attendance with optional geo/selfie", "CheckInOut.jsx → /both/attendance/clock-in|out", "Employee"),
        ("Geo-fencing", "Haversine vs attendance rules radius", "Client geolocation + attendance-rules", "Employee (rules: admin/hr)"),
        ("Selfie / camera capture", "Base64 image on attendance doc", "Clock-in UI canvas", "Employee"),
        ("Work-from-home punch", "Punch as work_from_home when org allows", "CheckInOut / AttendanceBox + allowWorkFromHome", "Employee"),
        ("WFH request (informational)", "Submit WFH dates; email notify; no approval chain", "WfhRequestPanel → /both/wfh-requests", "Any authenticated"),
        ("Today's attendance status", "Who is in/out/on leave", "TodayStatusComponent.jsx", "admin, hr, manager, executive"),
        ("Attendance history / analytics", "Historical views + WFH KPIs", "AllAttendanceStatsComponent.jsx", "admin, hr, manager, executive"),
        ("Attendance rules config", "Windows, geo, WFH, penalties, comp-off toggles", "AttendanceRules.jsx", "admin, hr, executive"),
        ("Attendance regularization", "Request + approve miss-punch fixes", "MyRegularizationRequests / RegularizationRequests", "Employee apply; manager/HR approve"),
        ("Biometric integration", "Fingerprint/face devices", "—", "—"),
        ("Attendance PDF export", "Export attendance PDF", "GET attendance/:employeeId/export-pdf", "Authenticated (scope risk)"),
    ],
    "Leave Management": [
        ("Apply / cancel leave", "Submit or cancel leave request", "LeaveApplication / hub → /employee/leave/*", "Any authenticated"),
        ("Approve / reject leave", "Manager/HR status update + balance + email", "LeaveManagement.jsx", "admin, hr, manager, executive"),
        ("Leave balance engine", "Accrual, carry-forward, pending, available", "services/leaveBalanceEngine.js", "System / HR adjust"),
        ("Leave policy CRUD", "Leave types, quotas, compensatory rules", "LeavePolicyManagement.jsx", "admin, hr, executive"),
        ("Admin balance viewer/adjust", "Inspect + PATCH balances", "EmployeeLeaveBalanceViewer + LeaveTab", "admin, hr, executive"),
        ("Leave reports", "Leave utilization reporting", "LeaveReportsPage.jsx", "admin, hr, executive"),
        ("Leave encashment", "Cash for unused leave", "LeaveEncashment.jsx + APIs (UI mostly unrouted)", "Employee apply; admin approve"),
        ("Comp-off (admin)", "Admin manage compensatory offs", "/api/v1/admin/comp-off", "admin, hr"),
        ("Comp-off (employee)", "Employee claim/use comp-off", "CompOffRoutes.js imported, never mounted", "Dead"),
        ("Approval workflows on leave", "Multi-step approval", "leave create → approval service", "Configured approvers"),
        ("Auto-escalation if no response", "Escalate if manager silent X days", "—", "—"),
    ],
    "Payroll & Compensation": [
        ("Salary structure CRUD", "Earnings/deduction components, PT", "SalaryStructureManager.jsx", "admin (API isAdminAuthenticated); UI also finance"),
        ("Generate / process payroll", "Bulk generate Payroll docs", "AdminPayroll.jsx → payrollController", "admin, hr, finance"),
        ("Tax / PF / ESI / TDS calc", "India-oriented statutory calculations", "payrollGenerationHelper, incomeTaxCalculation", "System"),
        ("Employee payslip + PDF", "View/download payslips", "EmployeePayroll.jsx + pdfExportController", "Self"),
        ("Bank disbursement / mark Paid", "Transfer salary + mark Paid", "paymentService.transferToBank → mock-bank", "CEO-gated disburse"),
        ("Legacy list all payrolls", "List org payrolls", "GET /admin/payroll/get-all-payroll (isAuthenticated only)", "Any authenticated — SECURITY GAP"),
        ("Payroll reports", "Payroll statistics", "PayrollReport.jsx", "admin, hr, finance, executive"),
    ],
    "Recruitment": [
        ("Recruitment ATS list", "List jobs + candidates", "RecruitmentPage.jsx", "admin, hr, executive"),
        ("Create job posting (API)", "POST job openings", "recruitmentController", "admin, hr, executive"),
        ("Create/detail job (UI)", "Buttons navigate to unregistered routes", "/recruitment/create-job — 404", "—"),
        ("Candidate apply", "Apply to job (authenticated)", "POST /both/recruitment/jobs/:id/apply", "Any authenticated"),
        ("Interviews / offer letters", "Backend endpoints; thin SPA", "Controller only", "admin, hr"),
        ("Recruitment report", "Funnel metrics (largely fake)", "RecruitmentReport.jsx", "admin, hr, executive"),
    ],
    "Performance & 1:1": [
        ("Employee performance tab", "Manual score + stats", "PerformanceTab.jsx", "admin, hr, manager"),
        ("Performance reviews / goals / cycles APIs", "CRUD review docs, goals, cycles", "/both/performance-review, /both/goal, /admin/review-cycle", "Mixed"),
        ("Goals / appraisal SPA", "Primary UI for cycles/goals", "No routes in routes/index.jsx", "—"),
        ("Performance report", "Rating distribution (fabricated fallbacks)", "PerformanceReport.jsx", "admin, hr, executive"),
        ("One-on-one meetings", "Manager ↔ employee 1:1", "OneOnOneRouteWrapper", "Any authenticated"),
    ],
    "Onboarding / Offboarding": [
        ("Onboarding workflows (HR)", "Create/manage checklists", "OnboardingPage.jsx → /admin/onboarding", "admin, hr, executive"),
        ("Onboarding detail", "Workflow tasks", "OnboardingDetailPage.jsx", "admin, hr, executive"),
        ("My onboarding (employee)", "Employee checklist", "MyOnboardingPage — route guards exclude employees", "Unreachable"),
        ("Offboarding / IT provisioning", "Exit checklist + access revoke", "Model hints only / not implemented", "—"),
    ],
    "Documents & ID Cards": [
        ("My documents", "Upload/download personal docs", "MyDocumentsPage.jsx", "Any authenticated"),
        ("HR documents approve/versions", "Review + versions", "DocumentsPage / DocumentVersions", "admin, hr, executive"),
        ("ID card generator", "Generate employee ID cards", "IdCardGenerator.jsx", "admin, hr, executive"),
        ("ID card test-db (debug)", "Dump sample employees", "GET /admin/id-cards/test-db — NO AUTH", "Public — SECURITY GAP"),
    ],
    "Policies & Approvals": [
        ("Company policy CRUD + acknowledgment", "Create/publish; employees acknowledge", "PolicyManager / ViewPolicies", "admin/hr create; all acknowledge"),
        ("Approval workflow config + pending inbox", "Multi-step workflows + act on items", "WorkflowManager / PendingApprovals", "admin, hr, manager, executive"),
    ],
    "Calendar": [
        ("Calendar management / employee view", "Holidays / events / working days", "CalendarManagement / EmployeeCalendar", "admin/hr manage; all view"),
        ("Calendar image OCR", "Extract holidays from image", "calendarUploadController — hardcoded stub", "admin"),
    ],
    "Reports & Analytics": [
        ("HR reports suite", "Headcount, attrition, attendance, payroll, performance, recruitment, compliance", "HRReportsLayout → /admin/reports/*", "admin, hr, executive"),
        ("HR / Org analytics", "Broader analytics dashboards", "HRAnalyticsPage / OrganizationAnalytics", "admin, hr, executive"),
        ("Excel export", "Export tables", "exportUtils.exportToExcel — actually CSV", "Callers"),
    ],
    "Notifications": [
        ("In-app + Socket.IO + web push", "Inbox, realtime, VAPID push", "Notifications.jsx + websocket + push-subscription", "Any authenticated"),
        ("Announcements + notification settings", "Compose org announcements; prefs", "NotificationSender / NotificationSettings", "admin, hr"),
        ("Transactional email / SMS", "Leave/auth emails; SMS flag only", "emailSender.js; no SMS provider", "System"),
    ],
    "Roles & Access Control": [
        ("Platform / system roles", "Assign fixed SaaS roles + permissions", "UsersRolesPage → /admin/platform-roles", "admin, hr"),
        ("Role hierarchy UI (custom roles)", "Create custom Role docs with parent/permissions", "RoleHierarchy.jsx → /admin/roles", "admin, hr, executive"),
        ("requirePermission on APIs", "Real RBAC on many routes", "middlewares/requirePermission.js", "Server"),
        ("Legacy Role.rbac middleware", "module/action Role model checks", "middlewares/rbac.js — unused by route mounts", "Cosmetic"),
        ("Admin Settings Permissions/Security tabs", "Placeholder theatre", "AdminSettings.jsx", "admin — cosmetic"),
    ],
    "Organization & Settings": [
        ("Org profile / settings / subscription / verification", "Identity, theme, billing, domain", "Admin/Organization/*", "admin, executive"),
        ("Email / theme settings", "SMTP + UI theme", "EmailSettings / ThemeSettings", "admin"),
        ("Audit logs", "Admin audit trail API", "/admin/audit-logs", "admin"),
    ],
    "Projects & Tasks": [
        ("Projects CRUD + employee tasks", "Manage projects / assigned work", "ProjectPage / EmployeeProjects / TaskRoutes", "admin, hr, manager, project_manager, executive; employees self"),
        ("Project delete", "Delete from list", "ProjectPage.jsx TODO — not wired", "—"),
    ],
    "Learning / Benefits / Engagement": [
        ("Learning hub + courses", "Catalog, create/edit, take course, team learning", "LearningPage / CreateCourse / CourseDetails / TeamLearning", "Any / HR create"),
        ("Benefits plans + enroll", "Manage/enroll benefit plans", "BenefitsPage / BenefitDetailPage", "admin, hr / any"),
        ("Engagement surveys + 360", "Surveys/feedback; 360 placeholder", "EngagementPage / CreateSurvey", "Any / admin create"),
        ("Training ROI", "ROI calculation", "learningController.calculateTrainingROI — hardcoded costs", "admin"),
    ],
    "Helpdesk & SLA": [
        ("Helpdesk ticket list", "List tickets", "HelpdeskPage.jsx — New/View dead; employees blocked", "admin, hr, executive"),
        ("SLA manager UI", "CRUD SLAs + violations", "SLAManager.jsx — NO ROUTE; API /violations shadowed", "Unreachable"),
    ],
    "Expenses / Succession / Workforce": [
        ("Expense API", "Submit/approve/mark-paid", "/both/expense", "Employee submit; admin approve"),
        ("Expense UI", "Finance dashboard → /admin/expenses", "No React route — 404", "—"),
        ("Succession / Workforce planning", "Plans + talent pool + forecasts", "SuccessionPage / WorkforcePage", "admin, hr, executive"),
    ],
    "Mobile / PWA": [
        ("Capacitor + responsive + PWA install", "Native shell + mobile layout + install prompt", "capacitor.js / Header / InstallPrompt", "All"),
        ("Manager mobile leave approve", "Responsive web; no offline queue", "Same SPA", "Managers"),
    ],
    "Security / Debug (side features)": [
        ("Health check", "Liveness", "/api/health", "Public"),
        ("Debug uploads listing", "List upload filesystem", "/api/debug/uploads — NO AUTH", "Public — SECURITY GAP"),
        ("Legacy unauth push subscribe", "POST /api/push/subscribe + send-test", "routes/push.js", "Public — SECURITY GAP"),
        ("Unauth profile photo GET/DELETE", "GET/DELETE /both/profile-photo/:type/:id", "ProfilePhotoRoutes.js", "Public — SECURITY GAP"),
    ],
}

# =============================================================================
# PHASE 1B — Role Permission Matrix
# =============================================================================
# Exists? Yes = first-class system/platform role; Partial = inferred from position/dept/functionalRoles;
# No = not in catalog.
ROLE_MATRIX = [
    ("CEO", "Partial", "executive bucket + CEO dashboard route; functionalRoles ceo/executive; EXECUTIVE_DASHBOARD, org-wide view perms", "No dedicated CEO permission set; no board reporting; no exclusive payroll disburse beyond UI gate; no company-wide policy veto", "Both (dashboard FE; limited BE executive routes)"),
    ("COO", "No", "None dedicated — would map only if titled/positioned as executive", "Ops KPIs, cross-dept approvals, business continuity — all missing as role", "N/A — frontend would need custom RoleHierarchy entry (not enforced by requirePermission)"),
    ("CTO", "Partial", "executive + TECH_DASHBOARD if dept=engineering/IT; functionalRoles cto/executive", "Tech hiring pipeline ownership, system access provisioning, engineering OKR cascade — missing", "Both (partial)"),
    ("CFO", "Partial", "executive + FINANCE_DASHBOARD if finance dept; functionalRoles cfo; PAYROLL_VIEW_ALL via executive", "Treasury, budget vs actual, bank file approval, statutory filings ownership — missing", "Both (partial)"),
    ("CHRO / CPO", "Partial", "Maps to hr_manager / HR_ADMIN / org_owner; full HR_MANAGER_PERMS", "People strategy OKRs, org design, compensation bands UI, board people report — thin", "Both"),
    ("CMO", "No", "None", "Marketing headcount/campaign staffing — not an HRMS role here", "N/A"),
    ("VP", "Partial", "functionalRoles vp/executive; VP dashboard path", "Delegation when absent; portfolio approvals — no delegation engine", "Both (partial)"),
    ("Director", "Partial", "Inferred as manager if jobTitle/position contains director", "Director-level multi-team span of control — same as line manager perms", "Both (partial)"),
    ("General Manager", "No", "No GM system role", "Site/P&L people authority — missing", "N/A"),
    ("Department Head", "Partial", "Treated as line_manager if position/jobTitle manager/director", "Dept budget headcount planning beyond workforce page — limited", "Both (partial)"),
    ("Team Lead", "Partial", "Only if titled manager; no distinct team_lead role", "Lead-only soft approve without full manager — missing", "Frontend role string only if inferred"),
    ("Manager / Line Manager", "Yes", "SYSTEM line_manager: team employee/leave/attendance view+approve, regularization, self payroll, project.admin.view", "Cannot formally see team salaries via permission matrix — but legacy get-all-payroll bypasses this; no skip-level; no out-of-office delegate", "Both (with BE gap on payroll list)"),
    ("Project Manager", "Partial", "functionalRoles project_manager; project routes", "Resource allocation, project timesheets, billing rates — thin", "Both (projects)"),
    ("HR Admin / HR Manager", "Yes", "hr_manager / hr_admin: nearly all PERMISSIONS except executive.dashboard", "CHRO board packs; multi-country policy packs — missing", "Both"),
    ("HR Executive", "Yes", "hr_executive: employee CRUD-ish, leave/attendance view/approve, ID cards; no payroll publish", "Letter generation, offer letters UI, background verification — partial/missing", "Both"),
    ("Recruiter / TA", "No", "Uses HR roles for /recruitment", "Dedicated recruiter scope (own reqs only), careers portal, interview scorecards — missing", "FE: HR only"),
    ("Payroll Officer", "Yes", "payroll_admin: payroll CRUD/process/publish + employee.view.all", "Bank file generation, statutory returns, payslip dispute — partial (calc OK, disburse mock)", "Both (salary structure API still isAdminAuthenticated)"),
    ("HR Business Partner", "No", "No HRBP role; would use HR executive/manager", "Client-group scoping, talent reviews — missing", "N/A"),
    ("Finance Manager", "Partial", "finance functional/dept → FINANCE_DASHBOARD; not a system role", "Expense approval UI missing; GL export — missing", "FE heavy; BE payroll gaps"),
    ("Accountant", "No", "None", "Journal posting, vendor payments — not in product", "N/A"),
    ("Payroll Finance Approver", "Partial", "Overlaps payroll_admin / CEO disburse UI gate", "Maker-checker dual control for payout — incomplete (mock bank)", "Partial BE"),
    ("System Admin / Org Owner", "Yes", "org_owner bucket = ALL_PERMISSIONS; org dashboard", "Tenant-level feature flags per industry; custom permission builder beyond RoleHierarchy (legacy unused) — partial", "Both"),
    ("IT Support", "Partial", "tech dashboard if IT dept; no IT helpdesk agent role", "Access provisioning tickets, device inventory — missing; helpdesk employee-blocked", "FE only mostly"),
    ("Regular Employee", "Yes", "employee: punch, leave apply, payroll self, leave.view.team", "Helpdesk create blocked; my-onboarding unreachable; encashment UI unrouted", "Both"),
    ("Intern", "No", "No intern employment-type role; employee only", "Limited leave types, stipend payroll, mentor assignment — missing", "N/A"),
    ("Contractor / Freelancer", "No", "No contractor type with restricted portal", "Limited self-service, no benefits, invoice-based pay — missing", "N/A"),
    ("Probation Employee", "No", "No probation flag-driven ACL", "Probation leave caps, confirmation workflow — missing as role (may exist as field elsewhere)", "N/A"),
    ("Auditor (read-only)", "No", "No auditor role; audit-logs admin-only", "Read-only compliance across all modules without mutate — missing", "N/A"),
    ("Legal / Compliance Officer", "No", "None; ComplianceReport is fake static UI", "Policy legal hold, POSH case mgmt, retention — missing", "N/A"),
]

ROLE_SYSTEM_NOTES = [
    "Permission matrix exists centrally in server/config/permissions.constants.js (PERMISSIONS + ROLE_PERMISSION_MAP) and is enforced on many APIs via requirePermission — this is real RBAC, not only UI.",
    "Parallel systems: (1) platform roles on User, (2) system roles (6 SaaS IDs), (3) functionalRoles on Employee inferred from dept/title, (4) custom Role model via RoleHierarchy — middlewares/rbac.js for Role docs is NOT wired into route mounts, so custom hierarchy permissions are largely cosmetic for API ACL.",
    "Custom roles: RoleHierarchy can CREATE custom roles for a client company, but Super Admin cannot reliably map those custom roles onto requirePermission keys used by APIs. UsersRolesPage only assigns the fixed 6 system roles.",
    "Multi-role: User.roles is an array; functionalRoles is an array — yes, one person can hold Manager + HR (e.g. functionalRoles includes both). Permission buckets are merged via permission service.",
    "Hierarchy/delegation: Role model has parentRole for UI tree; NO approve-on-behalf / out-of-office delegation engine found.",
    "CRITICAL: Several sensitive APIs use weak/no auth (get-all-payroll, profile-photo DELETE, push, test-db, debug uploads) — frontend hiding is irrelevant; see Section G.",
]

# =============================================================================
# PHASE 2 — Real-world gaps
# =============================================================================
GAPS = {
    "Attendance": [
        ("See who is on leave / in office today at a glance", "TodayStatusComponent + leave overlays", "Fully Addressed", "Keep leave+attendance joined loads"),
        ("Reliable geo/photo proof resisting spoofing", "Geo + selfie exist; client-spoofable; simple punches can skip", "Partially Addressed", "Server enforce + device attestation"),
        ("Biometric / factory devices", "Not implemented", "Not Addressed", "Vendor middleware"),
        ("WFH request with manager approval", "WFH request is informational email only; separate from punch", "Partially Addressed", "Wire approval workflow + link to attendance"),
        ("Managers approve regularization on mobile", "Responsive SPA; no offline queue", "Partially Addressed", "Push deep-links + QA"),
    ],
    "Leave": [
        ("Apply leave and track balance accurately", "Strong leaveBalanceEngine + apply/approve/cancel", "Fully Addressed", "Fail closed on engine errors"),
        ("Team leave calendar at a glance", "Partial via hub/today status", "Partially Addressed", "Manager team leave widget"),
        ("Auto-escalate if manager silent X days", "Workflows exist; no SLA auto-escalate job for leave", "Not Addressed", "Cron + SLA"),
        ("Comp-off earn/use by employees", "Admin API mounted; employee CompOffRoutes dead", "Partially Addressed", "Mount + UI"),
        ("Leave encashment", "API + orphaned UI", "Partially Addressed", "Route LeaveEncashment.jsx"),
    ],
    "Payroll": [
        ("Auto-calculate PF/ESI/TDS (India)", "Real engines with slabs/caps", "Fully Addressed", "CA sign-off + regression tests"),
        ("Multi-country tax engines", "India-oriented only", "Not Addressed", "Strategy per country"),
        ("Payslip download", "PDF path exists", "Fully Addressed", "Clarify publish workflow"),
        ("Actually disburse salaries", "mock-bank always succeeds; marks Paid", "Not Addressed", "Real payout provider"),
        ("Prevent non-payroll staff seeing salaries", "Legacy get-all-payroll isAuthenticated only", "Not Addressed", "PAYROLL_VIEW_ALL + scoping"),
        ("Expense reimbursements", "API only; UI 404", "Not Addressed", "Build Expenses page"),
    ],
    "Recruitment": [
        ("Post jobs and manage pipeline", "List API-backed; create/detail UI 404", "Partially Addressed", "Add create/detail pages"),
        ("Interview scheduling & offers", "Backend; thin SPA", "Partially Addressed", "Build UI"),
        ("Public careers portal", "Apply requires auth", "Not Addressed", "Guest careers page"),
        ("Trusted hiring analytics", "RecruitmentReport largely fake", "Not Addressed", "Bind to API; empty states"),
    ],
    "Performance": [
        ("Structured appraisal cycles", "APIs exist; no primary SPA", "Not Addressed", "Performance hub UI"),
        ("OKR / goal tracking", "Goal API only", "Not Addressed", "Goals UI"),
        ("Continuous 1:1s", "Routed + API-backed", "Fully Addressed", "Reminders/templates"),
        ("Calibration / 9-box", "Report fabricates matrix", "Not Addressed", "Persisted ratings only"),
    ],
    "Employee Lifecycle": [
        ("Onboarding checklist + docs", "HR workflows + documents", "Partially Addressed", "Open My Onboarding to employees"),
        ("IT access provisioning Day 1", "Manual Add Employee only", "Not Addressed", "IT checklist integrations"),
        ("Offboarding + access revoke", "Not a full exit product", "Not Addressed", "Dedicated offboarding"),
        ("Bulk hire CSV", "Missing", "Not Addressed", "Import with rollback"),
    ],
    "Compliance": [
        ("Statutory filing tracker", "ComplianceReport = hardcoded constants", "Not Addressed", "Compliance model + calendar"),
        ("Policy acknowledgment audit", "Policies + ack exist", "Partially Addressed", "Drive compliance dashboards from real data"),
        ("Audit trail of sensitive actions", "Audit log API partial coverage", "Partially Addressed", "Instrument all mutates; auditor role"),
    ],
    "Reporting / Analytics": [
        ("Trusted leadership reports", "Many HR report pages use fake charts", "Not Addressed", "Remove fake KPIs"),
        ("Real Excel export", "CSV mislabeled as Excel", "Partially Addressed", "ExcelJS"),
        ("Executive predictive insight", "Naive linear fit", "Partially Addressed", "Label illustrative or real models"),
    ],
    "Self-Service": [
        ("Update personal/bank/tax data", "ProfileDetails + tax + financials", "Fully Addressed", "Approve bank changes"),
        ("Documents vault", "My Documents + Documents&ID", "Fully Addressed", "Continue"),
        ("Raise helpdesk ticket", "CTA blocked; New Ticket dead", "Not Addressed", "Open + wire create"),
        ("View payslip/leave/attendance", "Yes", "Fully Addressed", "Continue"),
    ],
    "Notifications": [
        ("In-app approval alerts", "Notifications + Socket.IO", "Fully Addressed", "Prod socket URL QA"),
        ("Email for critical events", "Leave/auth; not exhaustive", "Partially Addressed", "Complete matrix"),
        ("SMS urgent alerts", "Not implemented", "Not Addressed", "Twilio/MSG91"),
        ("Mobile push for managers", "Web push partial", "Partially Addressed", "Cap Push + VAPID"),
    ],
    "Security / Access Control": [
        ("Backend-enforced RBAC", "requirePermission on many routes", "Partially Addressed", "Close legacy holes"),
        ("Settings UI matches ACL", "Permissions tab placeholder", "Not Addressed", "Deep-link to UsersRolesPage"),
        ("Protect debug endpoints", "test-db, debug uploads, unauth push open", "Not Addressed", "Remove/gate"),
        ("Hard delete preserves audit", "findOneAndDelete incomplete cascade", "Not Addressed", "Soft-delete + legal hold"),
    ],
}

# =============================================================================
# PHASE 2B — Industry readiness
# =============================================================================
INDUSTRY = [
    ("IT / Software", "Flexible shifts, remote/WFH, skill tags, project allocation, contractor eng", "Partial — WFH punch + WFH request + projects/tasks; no skill inventory; contractors=employees", "No tech-stack tagging; WFH lacks approval; no bench management", "Add skills model, WFH approval workflow, contractor employment type; keep projects module"),
    ("Manufacturing", "Shift attendance, biometric, OT/labor law, safety compliance", "No — no shift roster, no biometric, OT limited, no safety module", "Entire shop-floor stack missing", "Shift scheduler + device middleware + OT rules + safety incidents; do not sell as-is"),
    ("Retail", "Multi-location staff, hourly wages, roster/scheduling", "Partial — multi-org/tenant exists; no store roster or hourly timeclock product", "No location roster, no tip/hourly payroll modes", "Locations + shift templates + hourly payroll components"),
    ("Healthcare", "License/cert expiry alerts, on-call, HIPAA-like controls", "No dedicated clinical credentials; documents exist but no expiry alerts; no PHI controls", "Credentialing + on-call + stricter field encryption/access logs", "Certification tracker + encryption + audit for PHI; BAA/HIPAA program outside product"),
    ("BPO / Call centers", "Shift rotation, real-time attendance, AHT/productivity", "Attendance realtime-ish; no AHT/productivity; no shift rotation engine", "Workforce management (WFM) gap", "Integrate WFM or build shift + productivity metrics"),
    ("Construction / Field", "GPS attendance, site mapping, contractors", "Geo-fencing exists (spoofable); no sites model; no contractor portal", "Sites/projects geo zones + contractor ACL", "Site entities, geofence per site, contractor role"),
    ("Startups / SMB", "Simple approvals, minimal hierarchy", "Partial — can use employee+manager+HR; approval workflows configurable; RoleHierarchy overkill/confusing", "Too many half-built modules (reports fake, recruitment 404) hurt SMB trust", "Offer 'Essential' SKU: Auth, Employees, Attendance, Leave, basic Payroll; hide stub modules"),
]

INDUSTRY_FLEX = (
    "Permission/module structure is moderately flexible via system roles + RoleHierarchy UI, "
    "but API enforcement keys are hardcoded in permissions.constants.js. Industry packs cannot be "
    "toggled as feature flags per tenant today — modules are mostly always-on code paths. "
    "To sell as multi-industry SaaS: (1) tenant feature flags, (2) employment types, "
    "(3) custom role→permission mapping that actually hits requirePermission, "
    "(4) pluggable attendance/payroll strategies per industry."
)

# =============================================================================
# PHASE 3 — Status
# =============================================================================
STATUS_ROWS = {
    "Authentication & Session": [
        ("Unified login", OK, "Login.jsx → unifiedLogin → JWT cookies; bcrypt + lockout", "—"),
        ("Logout", OK, "clearAuthCookies", "—"),
        ("Organization signup/login", OK, "organizationAuthController", "Terms/privacy links may 404"),
        ("Admin register", PARTIAL, "Open when NODE_ENV !== production", "Misconfig risk"),
        ("Forgot / reset password", PARTIAL, "Hashed token path complete", "Needs SMTP — manual"),
        ("Session refresh", MANUAL, "Axios interceptor + refresh routes", "Runtime QA"),
        ("Frontend guards", OK, "ProtectedRoute + permissions", "—"),
        ("API middleware (overall)", PARTIAL, "Broad requirePermission coverage", "Critical unauth gaps remain"),
        ("MFA / SSO", BROKEN, "No TOTP/SAML/OAuth IdP", "Absent"),
    ],
    "Employee Management": [
        ("Add employee", OK, "Register → Employee.create; permission gated", "Rigid enums"),
        ("Employee list/search", PARTIAL, "Server regex; some filters client-only", "—"),
        ("Profile view / tabs", OK, "GET get-info + multi-tab UI", "—"),
        ("Soft terminate/status", OK, "PATCH terminate / update-status", "—"),
        ("Hard delete", BROKEN, "findOneAndDelete; payroll only cleaned", "Orphans + audit loss"),
        ("Bulk import/export", BROKEN, "No import; exportUtils is CSV", "Missing"),
        ("Profile photo DELETE", BROKEN, "No auth middleware on DELETE/GET", "Security defect"),
    ],
    "Attendance": [
        ("Clock in/out / breaks", OK, "Attendance.create/update server time", "Soft window warnings"),
        ("Geo-fencing", PARTIAL, "Haversine vs rules", "Client-spoofable"),
        ("Selfie capture", PARTIAL, "Base64 on doc", "Not biometric"),
        ("WFH punch + request", PARTIAL, "Punch + WfhRequestPanel + API; no approval", "Informational only"),
        ("Biometric", BROKEN, "No integration", "Absent"),
        ("Today status / regularization", OK, "APIs + UI; caps enforced", "—"),
        ("Holiday punch block", PARTIAL, "Calendar in analytics", "Punch not blocked"),
        ("Attendance PDF", MANUAL, "Route + isAuthenticated", "IDOR scope live test"),
    ],
    "Leave Management": [
        ("Apply / approve / cancel", OK, "Full chain + email + balance", "—"),
        ("Balance engine / policy CRUD", OK, "Engine + unit tests + UI", "—"),
        ("Admin balance adjust", OK, "LeaveTab PATCH", "—"),
        ("Leave encashment UI", BROKEN, "Component not in routes/index.jsx", "API may work"),
        ("Employee comp-off APIs", BROKEN, "compOffRouter never app.use'd", "Dead"),
        ("Auto-escalation", BROKEN, "No job", "Absent"),
        ("Soft-fail balance on apply", PARTIAL, "catch continues create", "Integrity risk"),
    ],
    "Payroll": [
        ("Generate + tax/PF/ESI", OK, "Real calculation helpers persist", "CA sign-off"),
        ("Employee payslip + PDF", OK, "History + generatePayrollPDF", "—"),
        ("Salary structure", PARTIAL, "CRUD for admin JWT", "Finance may 403"),
        ("Bank disbursement", BROKEN, "paymentService mock-bank always success", "Fake Paid"),
        ("get-all-payroll auth", BROKEN, "isAuthenticated only", "Salary leak"),
        ("Payroll report page", PARTIAL, "Mixed live/hardcoded", "—"),
    ],
    "Recruitment": [
        ("List jobs/candidates", PARTIAL, "RecruitmentPage fetches APIs", "Errors swallowed"),
        ("Create/detail job UI", BROKEN, "Unregistered routes → 404", "API orphaned"),
        ("Recruitment report", BROKEN, "Static/fake funnel", "—"),
    ],
    "Performance": [
        ("One-on-one", OK, "Routed pages + APIs", "—"),
        ("Performance tab score", PARTIAL, "Manual score only", "≠ appraisal"),
        ("Reviews/goals/cycles APIs", PARTIAL, "Mounted routers", "Manager auth often admin JWT"),
        ("Appraisal / goals UI", BROKEN, "No SPA routes", "Orphaned APIs"),
        ("Performance report", BROKEN, "Hardcoded distribution/KPIs", "Fabricated"),
    ],
    "Onboarding": [
        ("HR onboarding CRUD", PARTIAL, "Pages + /admin/onboarding", "—"),
        ("Employee my onboarding", BROKEN, "Role/perm gate excludes employees", "Unreachable"),
        ("IT provisioning / offboarding", BROKEN, "Not implemented", "—"),
    ],
    "Documents & ID": [
        ("Employee docs + HR approve", OK, "multer + Document model + admin APIs", "—"),
        ("ID card generator", OK, "IdCardGenerator + APIs", "—"),
        ("ID card test-db", BROKEN, "Unauthenticated PII dump", "Security"),
    ],
    "Reports": [
        ("Compliance report", BROKEN, "Static constants only", "No API"),
        ("Attendance / Attrition / Recruitment reports", BROKEN, "Hardcoded/fake series", "—"),
        ("Headcount report", PARTIAL, "Some live; FTE split invented", "—"),
        ("Excel export", BROKEN, "CSV mislabeled", "—"),
        ("Core dashboards", PARTIAL, "Real APIs often", "Errors → zeros"),
    ],
    "Notifications": [
        ("In-app + socket", OK, "Notification model + UI + socket.io", "—"),
        ("Web push", PARTIAL, "Auth'd subscription + legacy unauth /api/push", "VAPID"),
        ("Transactional email", PARTIAL, "emailSender substantial", "Incomplete matrix"),
        ("SMS", BROKEN, "Flag only", "—"),
    ],
    "Roles & Settings": [
        ("Platform roles API ACL", OK, "/admin/platform-roles + permissions", "—"),
        ("Custom RoleHierarchy", PARTIAL, "CRUD Role docs", "rbac.js unused by routes"),
        ("Settings Permissions/Security tabs", BROKEN, "Placeholders", "Theatre"),
    ],
    "Other modules": [
        ("Projects/tasks core", OK, "Admin + employee UIs + APIs", "Delete TODO"),
        ("Policies + approval workflows", OK, "CRUD + pending inbox", "—"),
        ("Calendar", PARTIAL, "Real holiday CRUD; OCR stub", "—"),
        ("Benefits / Learning / Engagement", PARTIAL, "Persist; ROI/360 stubs", "—"),
        ("Succession / Workforce", PARTIAL, "CRUD pages", "Depth limited"),
        ("Helpdesk", BROKEN, "List only; New dead; employees blocked", "—"),
        ("SLA UI", BROKEN, "Unrouted; /violations shadowed", "—"),
        ("Expenses UI", BROKEN, "API-only; Finance link 404", "—"),
        ("Assets / shifts module", BROKEN, "Not in codebase", "Absent"),
    ],
}

# =============================================================================
# PHASE 3B — Fixes
# =============================================================================
FIXES = [
    ("Payroll", "Bank disbursement / mark Paid",
     "paymentService.transferToBank always returns success with provider mock-bank",
     "1. Real payout adapter. 2. Mock behind PAYOUT_MODE. 3. Never Paid unless provider confirms. 4. Reconciliation job.",
     "Strategy interface {transfer,getStatus}; on mock keep DisbursementPending + SIMULATED banner",
     "Large", "Yes — payroll status, payslips, finance reports"),
    ("Payroll", "GET get-all-payroll auth hole",
     "PayrollRoutes uses isAuthenticated only",
     "1. requirePermission(PAYROLL_VIEW_ALL) or scoped handler. 2. Org+team filter. 3. Employee 403 test.",
     "router.route('/get-all-payroll').get(requirePermission(PERMISSIONS.PAYROLL_VIEW_ALL), getAllPayrolls)",
     "Quick", "No — tightens only"),
    ("Payroll", "Salary structure finance 403",
     "isAdminAuthenticated vs finance UI",
     "1. Align to payroll RBAC permissions.",
     "composeMiddleware with PAYROLL_EDIT",
     "Medium", "Possibly"),
    ("Employee", "Hard delete orphans",
     "findOneAndDelete; payroll only cleaned",
     "1. Soft-delete. 2. Or cascade anonymize. 3. Legal hold block.",
     "status:'deleted', deletedAt:new Date()",
     "Medium", "Yes — leave/attendance/auth"),
    ("Employee", "Bulk import/export",
     "No import; CSV as Excel",
     "1. POST bulk-import. 2. Validate. 3. ExcelJS export.",
     "Streaming CSV + batch insert",
     "Large", "Yes — register path"),
    ("Leave", "Soft-fail balance on apply",
     "catch balanceError continues",
     "1. Return 503/409. 2. Do not create. 3. Test.",
     "return sendError(res, 503, 'LEAVE_BALANCE_UNAVAILABLE', ...)",
     "Quick", "Yes — can block apply if engine down"),
    ("Leave", "CompOffRoutes unmounted",
     "import without app.use",
     "1. app.use('/api/v1/both/comp-off', compOffRouter). 2. SPA tabs.",
     "app.use('/api/v1/both/comp-off', compOffRouter);",
     "Medium", "Yes — leave balances"),
    ("Leave", "Encashment UI unrouted",
     "LeaveEncashment.jsx orphaned",
     "1. Register route. 2. Link from hub.",
     "<Route path='/leave-encashment' .../>",
     "Quick", "Touches payroll if paid"),
    ("Leave", "No auto-escalation",
     "No SLA cron for leave",
     "1. Cron pending > SLA. 2. Escalate/notify.",
     "Reuse SLA + node-cron",
     "Large", "Yes — approvals/notifications"),
    ("Attendance", "WFH no approval",
     "Controller is informational email only",
     "1. Status pending/approved. 2. Manager approve API. 3. Optional auto-allow punch.",
     "Extend wfhRequest.model + approvalWorkflow",
     "Medium", "Attendance rules"),
    ("Attendance", "Client-spoofable geo",
     "Coords from client",
     "1. Reject if enforceGeo and invalid. 2. Size-limit selfie.",
     "Validate lat/lng when rule.enforceGeo",
     "Medium", "Rules"),
    ("Recruitment", "Create/detail 404",
     "Routes never registered",
     "1. JobForm + Detail pages. 2. Register routes. 3. Wire APIs.",
     "CRUD forms → /admin/recruitment/jobs",
     "Large", "Reports"),
    ("Recruitment", "Report fake",
     "Static funnel rows",
     "1. Call analytics API. 2. Empty state. 3. Remove hardcodes.",
     "useQuery recruitment-analytics",
     "Medium", "No"),
    ("Performance", "No appraisal UI",
     "APIs orphaned",
     "1. Performance hub. 2. Cycles/goals routes. 3. Fix manager auth.",
     "/performance/* route group",
     "Large", "Yes — analytics"),
    ("Performance", "Report fabricated",
     "Hardcoded [2,8,24,12,4]",
     "1. Real aggregate. 2. EmptyState. 3. Remove matrix fabrication.",
     "if (!ratings.length) return <EmptyState/>",
     "Quick", "No"),
    ("Reports", "ComplianceReport hardcoded",
     "No fetch",
     "1. Compliance model/API or remove page.",
     "ComplianceItem schema",
     "Large", "Compliance"),
    ("Reports", "Attendance/Attrition fake",
     "Invented cells/series",
     "1. Bind summary APIs. 2. Empty states.",
     "Map attendance.summary.*",
     "Medium", "No"),
    ("Reports", "exportToExcel is CSV",
     "Wrong format",
     "1. exceljs. 2. Separate exportCSV.",
     "workbook.xlsx.writeBuffer()",
     "Medium", "All exporters"),
    ("Onboarding", "My Onboarding unreachable",
     "allowedRoles exclude employees",
     "1. Fix routes/index.jsx employee path.",
     "Separate /onboarding/my for employees",
     "Quick", "No"),
    ("Helpdesk", "Dead New Ticket + employee blocked",
     "No onClick; role gate",
     "1. Create modal. 2. Detail route. 3. Allow employees own tickets.",
     "navigate(`/helpdesk/${id}`)",
     "Medium", "SLA"),
    ("SLA", "UI unrouted; /violations shadowed",
     "No SPA route; route order wrong",
     "1. Register /admin/sla. 2. /violations before /:id.",
     "router.get('/violations', ...)",
     "Quick", "Helpdesk"),
    ("Expenses", "UI 404",
     "No page; Finance links dead",
     "1. ExpensesPage. 2. Route. 3. Wire /both/expense.",
     "List + approve like leave",
     "Large", "Payroll reimbursements"),
    ("Security", "Profile photo DELETE unauth",
     "No middleware",
     "1. Auth. 2. Self or admin only.",
     "router.delete(..., isAuthenticatedBoth, deleteProfilePhoto)",
     "Quick", "No"),
    ("Security", "ID cards test-db / debug uploads / push",
     "Unauthenticated diagnostics",
     "1. Remove in prod. 2. Else admin+non-prod.",
     "if (NODE_ENV==='production') disable",
     "Quick", "No"),
    ("Security", "Public /uploads static",
     "app.use('/uploads') without auth",
     "1. Signed URLs. 2. Or auth middleware for sensitive docs.",
     "Stream via authenticated download route only",
     "Medium", "Documents"),
    ("Admin Settings", "Permissions stubs",
     "Placeholder; Manage Roles no-op",
     "1. navigate('/admin/platform-roles'). 2. Remove theatre.",
     "onClick={() => navigate('/admin/platform-roles')}",
     "Quick", "No"),
    ("Roles", "Dual RBAC illusion",
     "RoleHierarchy vs unused rbac.js",
     "1. Document platform roles as SoT. 2. Wire or deprecate hierarchy toggles.",
     "Map Role toggles → PERMISSIONS or remove",
     "Medium", "Yes — all ACL"),
    ("Calendar", "OCR hardcoded",
     "getCommonHolidays stub",
     "1. Real OCR or 501. 2. Label sample.",
     "Return 501 until Vision API",
     "Medium", "Working days"),
    ("Learning", "Training ROI hardcoded",
     "₹5000 / ₹50 magic numbers",
     "1. Store costs. 2. Or remove widget.",
     "course.cost * enrollments",
     "Quick", "No"),
    ("Auth", "sessionStorage org admin password",
     "plaintext password for Add Employee autofill",
     "1. Stop storing. 2. Invite tokens.",
     "Remove sessionStorage password",
     "Quick", "Add employee UX"),
    ("Projects", "Delete not wired",
     "TODO in ProjectPage",
     "1. DELETE API. 2. Confirm modal.",
     "axios.delete(`/v1/admin/project/${id}`)",
     "Quick", "Tasks cascade?"),
]

# =============================================================================
# PHASE 3C — Security
# =============================================================================
SECURITY = [
    ("Authentication", "Password strength policy", "⚠️", "Medium",
     "inputValidation.js enforces 8+ with complexity on many auth routes; user.model only minlength 8; reset path weaker",
     "Enforce same validator on all password set/reset paths; add zxcvbn"),
    ("Authentication", "Password expiry / reuse prevention", "❌", "Medium",
     "No expiry or password history found",
     "Store password hashes history; force rotate policy for admins"),
    ("Authentication", "MFA / 2FA", "❌", "High",
     "No TOTP/SMS/authenticator in codebase",
     "Add TOTP for org_owner/payroll_admin at minimum"),
    ("Authentication", "Session management (JWT expiry/refresh)", "⚠️", "Medium",
     "Access ~1h, refresh ~30d; passwordChangedAt invalidation; no concurrent session limit; same JWT_SECRET for both",
     "Separate refresh secret; session store/device limit; shorter refresh for privileged roles"),
    ("Authentication", "SSO (SAML/OAuth)", "❌", "High",
     "Local email/password only",
     "Enterprise IdP (SAML/OIDC) for SaaS sales"),
    ("Authorization", "Backend RBAC enforcement", "⚠️", "High",
     "requirePermission + ROLE_PERMISSION_MAP real on many routes; legacy holes (get-all-payroll, photos, push, debug)",
     "Inventory all routes; deny-by-default; close listed holes"),
    ("Authorization", "Frontend-only hiding", "⚠️", "High",
     "ProtectedRoute real but insufficient alone — several APIs lack matching checks",
     "Never rely on UI; add integration authz tests"),
    ("Authorization", "IDOR / cross-tenant access", "⚠️", "High",
     "org context middleware sets orgId but not all queries filter; work-info findById without organizationId; test-db global Employee.find",
     "Mandate organizationId on every tenant query; IDOR test suite"),
    ("Authorization", "Field-level salary restriction", "❌", "High",
     "Managers lack PAYROLL_VIEW_ALL in matrix but get-all-payroll lets any authed user see visible payrolls + salary populate",
     "Fix route auth; strip salary from populate for unauthorized"),
    ("Data Protection", "Encryption at rest (bank/Aadhaar/salary)", "❌", "High",
     "financialDetails, nationalId, salary stored plaintext in Mongo",
     "Field-level encryption (KMS); tokenize account numbers"),
    ("Data Protection", "Encryption in transit (HTTPS)", "⚠️", "Medium",
     "App does not enforce HTTPS itself; depends on reverse proxy; cookies need Secure in prod",
     "HSTS at edge; reject non-TLS in production config"),
    ("Data Protection", "File upload security", "⚠️", "High",
     "Documents have fileFilter; /uploads served statically without auth; profile MIME-only",
     "Authenticated download only; malware scan; block public static for PII"),
    ("Data Protection", "Data masking in logs/views", "⚠️", "Medium",
     "logger exists; no systematic PAN/Aadhaar masking verified",
     "Redact PII in queryLogger/email logs"),
    ("Multi-Tenant", "Tenant isolation", "⚠️", "High",
     "Shared DB + organizationId filters (application-level); not Mongo RLS; bugs can leak",
     "Helper that throws if org filter missing; periodic audit of findById"),
    ("Application", "SQL/NoSQL injection", "✅", "Low",
     "Mongoose ORM + mongo-sanitize middleware",
     "Continue; avoid $where / raw operators from user input"),
    ("Application", "XSS protection", "✅", "Low",
     "xss middleware + Helmet; React escapes by default",
     "Keep DOMPurify for any rich text"),
    ("Application", "CSRF protection", "✅", "Low",
     "Global CSRF middleware + token store",
     "Ensure SPA always sends token; verify cookie SameSite"),
    ("Application", "Rate limiting (login/sensitive)", "✅", "Low",
     "loginLimiter + apiLimiter; skipped for /uploads",
     "Add stricter limits on password reset and push"),
    ("Application", "Audit logs", "⚠️", "Medium",
     "auditLog model + admin viewer; used in leave/expense/payroll/employeeManagement/attendanceRule — not auth, docs, WFH, photos, recruitment",
     "Instrument all sensitive mutates; immutable storage"),
    ("Compliance", "GDPR export / right to be forgotten", "❌", "High",
     "Org gdprConsent flags only; no subject export/erasure APIs",
     "Data export package + anonymize cascade"),
    ("Compliance", "Local labor law packs", "⚠️", "Medium",
     "India PF/ESI/TDS engines; not multi-jurisdiction",
     "Country strategy modules"),
    ("Compliance", "Backup / DR process", "❌", "High",
     "Not in application codebase (ops concern)",
     "Document Mongo backups, RPO/RTO, restore drills"),
]

CRITICAL = [
    ("P0 — Money", "Mock bank marks payroll Paid", "Operators believe salaries disbursed; paymentService.js always success."),
    ("P0 — Security", "GET /admin/payroll/get-all-payroll under-authed", "Any logged-in employee may see org payroll + salary fields."),
    ("P0 — Security", "Unauthenticated profile-photo GET/DELETE", "Anyone can fetch/delete avatars by ID. ProfilePhotoRoutes.js."),
    ("P0 — Security", "Unauthenticated /api/push subscribe + send-test", "Push spam/hijack. routes/push.js."),
    ("P0 — Security", "Unauthenticated id-cards/test-db + /api/debug/uploads", "PII sample dumps + filesystem listing."),
    ("P0 — Security", "Public /uploads static serving", "Guessable filenames expose documents/photos without auth."),
    ("P0 — Data integrity", "Hard employee delete without cascade", "Orphans leave/attendance/docs; destroys audit trail."),
    ("P0 — Decision risk", "HR Reports fabricate compliance/attendance/performance/attrition/recruitment numbers", "Leadership can act on fake gauges."),
    ("P1 — Leave integrity", "Leave apply soft-fails if balance engine throws", "Leave created without balance enforcement."),
    ("P1 — Broken paths", "Recruitment create/detail, Expenses, SLA UI, employee onboarding, helpdesk create", "404/no-op product promises."),
    ("P1 — Comp-off", "Employee CompOff router imported never mounted", "Half-built feature."),
    ("P1 — ACL theatre", "Admin Settings Permissions + unused rbac.js vs real platform roles", "False sense of custom RBAC."),
    ("P1 — Field encryption", "Bank/PAN/Aadhaar/salary plaintext at rest", "SaaS breach impact High."),
    ("P2 — Secrets hygiene", "Plaintext password in sessionStorage (orgAdminCredentials)", "XSS exposure."),
]

ROADMAP = [
    ("P0", "Security", "Lock down get-all-payroll, photo DELETE/GET, push, test-db, debug, public /uploads", "Hours"),
    ("P0", "Payroll", "Stop marking Paid on mock bank; label simulated payouts", "Hours–days"),
    ("P0", "Reports", "Strip/disable fake Compliance/Attendance/Performance/Attrition/Recruitment KPIs", "Hours"),
    ("P0", "Employee", "Replace hard delete with soft-delete", "Hours"),
    ("P1", "Leave", "Fail closed on balance errors; mount CompOffRoutes; route encashment", "Hours"),
    ("P1", "Attendance", "WFH approval workflow; stronger geo enforce", "Days"),
    ("P1", "Recruitment", "Ship create-job + detail pages", "Days"),
    ("P1", "Helpdesk / Expenses / SLA", "Wire UI routes + handlers; employee tickets", "Days"),
    ("P1", "Onboarding", "Open My Onboarding to employees", "Hours"),
    ("P1", "Auth hygiene", "Remove sessionStorage password; gate admin register", "Hours"),
    ("P1", "Roles", "Document SoT ACL; deprecate unused rbac or wire it; add Auditor read-only role", "Days"),
    ("P2", "Security", "MFA, SSO, field encryption, GDPR export/erasure, session limits", "Weeks"),
    ("P2", "Performance", "Appraisal/goals SPA; fix manager auth", "Weeks"),
    ("P2", "Payroll", "Real payout provider + reconciliation", "Weeks"),
    ("P2", "Industry", "Feature flags + employment types + shift packs for multi-industry SaaS", "Weeks–months"),
    ("P2", "Leave", "Auto-escalation / SLA", "Days–weeks"),
    ("P2", "Compliance", "Real statutory filings tracker", "Weeks"),
    ("P2", "Notifications", "SMS + complete email matrix + Cap push", "Weeks"),
]


def count_statuses():
    counts = {OK: 0, PARTIAL: 0, BROKEN: 0, MANUAL: 0}
    for rows in STATUS_ROWS.values():
        for _, st, _, _ in rows:
            counts[st] = counts.get(st, 0) + 1
    total = sum(counts.values())
    return counts, total


def count_security():
    ok = sum(1 for r in SECURITY if r[2] == "✅")
    partial = sum(1 for r in SECURITY if r[2] == "⚠️")
    missing = sum(1 for r in SECURITY if r[2] == "❌")
    high = sum(1 for r in SECURITY if r[3] == "High" and r[2] != "✅")
    return ok, partial, missing, high


def build():
    counts, total = count_statuses()
    fully = counts[OK]
    partial = counts[PARTIAL] + counts[MANUAL]
    broken = counts[BROKEN]
    completion = round(100 * (fully + 0.5 * partial) / total) if total else 0
    sec_ok, sec_partial, sec_missing, sec_high = count_security()

    inv_count = sum(len(v) for v in INVENTORY.values())
    roles_yes = sum(1 for r in ROLE_MATRIX if r[1] == "Yes")
    roles_partial = sum(1 for r in ROLE_MATRIX if r[1] == "Partial")
    roles_no = sum(1 for r in ROLE_MATRIX if r[1] == "No")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    # ----- Cover -----
    add_para(doc, "HRMS Functionality, Role, Industry & Security Audit Report",
             size=22, bold=True, center=True, color=RGBColor(0x1A, 0x36, 0x5D), space_after=12)
    add_para(doc, f"Date: {datetime.now().strftime('%B %d, %Y')}", size=12, center=True)
    add_para(doc, "Project: EMS (Employee Management System) — server + vite_client", size=12, center=True)
    add_para(doc, "Method: Static code-path tracing (UI → route → controller → model/DB). "
             "No live HTTP/SMTP/bank runs. Runtime-dependent items marked Needs manual testing.",
             size=10, center=True, color=RGBColor(0x55, 0x55, 0x55), space_after=12)
    add_para(doc, "Classification: Internal — Stakeholder / Management", size=10, center=True, bold=True)
    add_para(doc, "Phases covered: 1 Inventory · 1B Roles · 2 Real-world gaps · 2B Industry · "
             "3 Working status · 3B Fixes · 3C Security", size=9, center=True, space_after=18)

    doc.add_page_break()

    # ----- Executive Summary -----
    doc.add_heading("Executive Summary", level=1)
    add_para(doc,
             f"Overall estimated completion for a full-suite multi-industry HRMS: ~{completion}% "
             f"(feature census n={total}: ✅ {fully} Fully Working, ⚠️ {partial} Partial/Manual, ❌ {broken} Not Working). "
             f"Inventory lists {inv_count} features across {len(INVENTORY)} modules. "
             f"Of {len(ROLE_MATRIX)} requested real-world roles: Yes {roles_yes}, Partial {roles_partial}, No {roles_no}. "
             f"Security checks: ✅ {sec_ok} / ⚠️ {sec_partial} / ❌ {sec_missing} ({sec_high} High-risk gaps). "
             "Core spine (auth, employee CRUD, leave engine, attendance punch/regularization, India payroll calc/PDF) is real. "
             "Outer suite (trusted HR reports, bank payout, recruitment UX, appraisal UI, expenses, MFA/SSO, field encryption) is incomplete or unsafe.")

    doc.add_heading("Status counts", level=2)
    add_table(doc, ["Status", "Count", "Share"], [
        (sym(OK), str(fully), f"{round(100 * fully / total)}%"),
        ("⚠️ Partially Working / Needs manual testing", str(partial), f"{round(100 * partial / total)}%"),
        (sym(BROKEN), str(broken), f"{round(100 * broken / total)}%"),
    ], status_col=0)

    doc.add_heading("Top 5 critical gaps (real-world needs unmet)", level=2)
    for i, g in enumerate([
        "Actual bank salary disbursement (mock-bank only) and multi-country payroll.",
        "Trusted leadership reports (compliance/attendance/performance/attrition/recruitment often fabricated).",
        "End-to-end recruitment UX (create/detail 404), expenses UI (404), appraisal SPA (APIs orphaned).",
        "Enterprise access: MFA, SSO, auditor/contractor/intern roles, true custom-role API enforcement.",
        "Employee self-service completeness: helpdesk create blocked, my-onboarding unreachable, WFH without approval.",
    ], 1):
        doc.add_paragraph(f"{i}. {g}", style="List Number")

    doc.add_heading("Top 5 critical bugs", level=2)
    for i, b in enumerate([
        "Mock payout marks payroll Paid with no money movement.",
        "GET /api/v1/admin/payroll/get-all-payroll auth too weak (salary data leak).",
        "Hard employee delete without cascade / soft-delete.",
        "Leave balance check soft-fails (leave still created if engine throws).",
        "CompOff employee router imported but never mounted; recruitment/SLA/expenses/encashment UI disconnected.",
    ], 1):
        doc.add_paragraph(f"{i}. {b}", style="List Number")

    doc.add_heading("Top 5 security risks", level=2)
    for i, s in enumerate([
        "Unauthenticated profile-photo DELETE/GET, /api/push subscribe/send-test, id-cards/test-db, debug uploads.",
        "Public /uploads static serving of documents and photos.",
        "Payroll salary exposure via under-authed get-all-payroll (IDOR/field leak).",
        "No MFA/SSO; plaintext bank/PAN/Aadhaar/salary at rest; incomplete org-scoped queries.",
        "No GDPR subject export/erasure; sparse audit coverage on sensitive mutates.",
    ], 1):
        doc.add_paragraph(f"{i}. {s}", style="List Number")

    add_para(doc,
             "Verdict: Do not sell or go-live as a complete multi-industry HRMS until P0 security/payroll/report issues are fixed. "
             "Leave / Attendance / India Payroll-calc / Auth are closest to production-grade. "
             "Permission enforcement is BOTH frontend and backend for core modules — but critical API holes make 'backend enforced' incomplete.",
             bold=True, space_after=12)

    doc.add_page_break()

    # ----- Section A -----
    doc.add_heading("Section A: Full Functionality Inventory", level=1)
    add_para(doc, "Every feature discovered in frontend + backend + models, grouped by module. "
             "Presence of a file does not imply the feature works (see Section E).")
    for module, rows in INVENTORY.items():
        doc.add_heading(module, level=2)
        add_table(doc, ["Feature", "Description", "Location", "Access Role"], rows)

    doc.add_page_break()

    # ----- Section B -----
    doc.add_heading("Section B: Role & Permission Matrix", level=1)
    add_para(doc, "Source of truth for API ACL: server/config/permissions.constants.js + requirePermission. "
             "System roles: employee, hr_executive, hr_manager, payroll_admin, line_manager, executive. "
             "Platform roles: org_owner, hr_admin, hr_executive, payroll_admin, manager, employee.")
    for note in ROLE_SYSTEM_NOTES:
        doc.add_paragraph(note, style="List Bullet")

    doc.add_heading("Role matrix (requested real-world roles)", level=2)
    add_table(doc,
              ["Role", "Exists?", "Modules Accessible", "Missing Real-World Permissions", "Enforced Where"],
              ROLE_MATRIX,
              status_col=1,
              status_map=YES_FILL)

    doc.add_heading("Permission matrix clarity", level=2)
    add_para(doc,
             "Yes — a clear role→permission map exists in permissions.constants.js (ROLE_PERMISSION_MAP). "
             "It is NOT only scattered string checks, though some legacy routes still use isAdminAuthenticated / isAuthenticated without PERMISSIONS keys. "
             "Custom roles via RoleHierarchy are NOT reliably enforced on APIs (rbac.js unused). "
             "Multi-role support: Yes (User.roles[] + functionalRoles[]). "
             "Delegation (approve on behalf of absent VP): No.")

    doc.add_page_break()

    # ----- Section C -----
    doc.add_heading("Section C: Real-World Gap Analysis", level=1)
    add_para(doc, "What HR teams need vs what this codebase actually delivers.")
    for module, rows in GAPS.items():
        doc.add_heading(module, level=2)
        add_table(doc,
                  ["Real-World Need", "Current Solution", "Status", "Recommendation"],
                  rows,
                  status_col=2,
                  status_map=GAP_FILL)

    doc.add_heading("Half-solutions & over-engineering", level=2)
    add_para(doc, "Half-solutions: WFH request without approval; comp-off admin-only; recruitment list without create UI; "
             "approval workflows without leave auto-escalation; RoleHierarchy without API enforcement; "
             "web push alongside unauthenticated legacy push.")
    add_para(doc, "Over-engineering relative to reliability: Large HR Reports/Analytics surfaces that invent KPIs; "
             "SLA/Helpdesk/Succession/Workforce UIs ahead of core security and payroll honesty; "
             "OCR calendar import that returns hardcoded holidays.")

    doc.add_page_break()

    # ----- Section D -----
    doc.add_heading("Section D: Industry Readiness Analysis", level=1)
    add_para(doc, INDUSTRY_FLEX)
    add_table(doc,
              ["Industry", "Required Roles/Workflows", "Currently Supported?", "Gap", "Recommendation"],
              INDUSTRY,
              status_col=2)

    doc.add_page_break()

    # ----- Section E -----
    doc.add_heading("Section E: Working / Not Working Status", level=1)
    add_para(doc, "Status key: ✅ Fully Working | ⚠️ Partially Working | ❌ Not Working | ⚠️ Needs manual testing. "
             "Brutally honest code-path verification — file existence ≠ working.")
    for module, rows in STATUS_ROWS.items():
        doc.add_heading(module, level=2)
        formatted = [(f, sym(st), e, iss) for f, st, e, iss in rows]
        add_table(doc, ["Feature", "Status", "Evidence", "Issue"], formatted, status_col=1)

    doc.add_page_break()

    # ----- Section F -----
    doc.add_heading("Section F: How to Fix — Solutions for Broken/Partial Features", level=1)
    add_para(doc, "Root cause + fix steps + code/approach + effort + cross-module impact.")
    add_table(doc,
              ["Module", "Feature", "Root Cause", "Fix Steps", "Suggested Code/Approach", "Effort", "Affects Other Modules?"],
              FIXES)

    doc.add_page_break()

    # ----- Section G -----
    doc.add_heading("Section G: Security Audit Findings", level=1)
    add_para(doc, "Traced in code (not assumed from config alone). Cross-references Phase 1B enforcement findings. "
             "Risk: High = exploit/data breach/money impact; Medium = policy/compliance gap; Low = mitigated controls present.")
    add_table(doc,
              ["Category", "Check", "Status", "Risk Level", "Evidence", "Recommendation"],
              SECURITY,
              status_col=2,
              status_map=SEC_FILL)

    doc.add_heading("Backend vs frontend enforcement (explicit)", level=2)
    add_para(doc,
             "Core HR modules (leave approve, employee create, payroll process/publish, platform roles) enforce on BOTH "
             "frontend (ProtectedRoute + requiredPermissions) and backend (requirePermission). "
             "However, several high-impact endpoints are Frontend-irrelevant because they lack backend auth entirely or use "
             "isAuthenticated only — this is a security issue, not a UX gap. Treat as P0 before any external tenant data.")

    doc.add_page_break()

    # ----- Section H -----
    doc.add_heading("Section H: Critical Issues", level=1)
    add_para(doc, "Bugs that can break production, leak/lose data, or mislead management — pulled from Sections E and G, sorted by severity.")
    add_table(doc, ["Severity", "Issue", "Detail"], CRITICAL)

    doc.add_heading("Dummy / mock data still in use", level=2)
    add_table(doc, ["Location", "What"], [
        ("server/helpers/paymentService.js", "Always-success mock-bank + MOCK-* references"),
        ("vite_client/.../ComplianceReport.jsx", "Entire report = static constants (2023 sample dates)"),
        ("AttendanceReport.jsx", "Hardcoded weeks; present 22/22; invented leave/late cells"),
        ("PerformanceReport.jsx", "Default rating 4.2; hardcoded distribution; fake matrix"),
        ("AttritionReport.jsx", "Chart fallbacks + fabricated series"),
        ("RecruitmentReport.jsx", "Static funnel / fake job rows"),
        ("HeadcountReport.jsx", "Invented ~85%/15% FTE/contract split"),
        ("calendarUploadController.js", "OCR returns hardcoded Indian holidays"),
        ("learningController.calculateTrainingROI", "Hardcoded ₹5000 / ₹50 economics"),
        ("engagementController 360", "Placeholder aggregation"),
        ("exportUtils.exportToExcel", "CSV, not Excel"),
        ("AdminSettings.jsx", "Security/Permissions/Reports placeholder tabs"),
        ("Employee Dashboard Todo.jsx", "localStorage-only todos (not HRMS data)"),
        ("wfhRequestController", "Requests are informational — no approval state machine"),
    ])

    doc.add_heading("Broken wiring (exists but disconnected)", level=2)
    add_table(doc, ["Item", "Problem"], [
        ("compOffRouter", "Imported in config/routes.js; never app.use'd"),
        ("/recruitment/create-job, jobs/:id, candidates/:id", "No React routes → 404"),
        ("/admin/expenses", "Finance dashboard link → 404"),
        ("SLAManager.jsx", "Never registered in routes/index.jsx"),
        ("SLA GET /violations", "Shadowed by /:id route order"),
        ("LeaveEncashment.jsx", "No SPA route"),
        ("Goals / performance-review / review-cycle", "APIs live; no primary SPA"),
        ("Helpdesk New Ticket / View", "Buttons without handlers"),
        ("MyOnboardingPage", "Route guards exclude employees"),
        ("middlewares/rbac.js", "Not imported by route mounts"),
        ("WFH requests", "API+UI exist; no approval / no link to leave workflows"),
    ])

    doc.add_page_break()

    # ----- Section I -----
    doc.add_heading("Section I: Recommendations & Priority Roadmap", level=1)
    add_para(doc, "Fix order: P0 before any production traffic with real salaries or external tenants; "
             "P1 to make advertised modules usable; P2 for enterprise/multi-industry depth.")
    add_table(doc, ["Priority", "Module", "Action", "Effort band"], ROADMAP)

    doc.add_heading("Recommended near-term sequencing", level=2)
    for t in [
        "Week 1: P0 security patches (auth holes, public uploads) + stop mock Paid + strip/flag fake report pages.",
        "Week 2: Leave fail-closed + mount comp-off + employee onboarding route + Settings stubs → real pages + soft-delete.",
        "Week 3–4: Recruitment create/detail UI; Expenses list UI; Helpdesk create/view; SLA route fix; WFH approval.",
        "Month 2: MFA for privileged roles; field encryption plan; real payout provider; Performance SPA; compliance tracker.",
        "Month 3+: SSO, GDPR export/erasure, industry feature flags, auditor role, shift/biometric packs as SKUs.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("What this audit did not prove", level=2)
    add_para(doc,
             "Live Mongo under concurrent tenants, SMTP delivery in your environment, Capacitor store builds, "
             "real-office geofence accuracy, CA-signed tax compliance, and adversarial org-id isolation all need "
             "staging/manual/penetration tests before go-live sign-off.")

    doc.add_heading("Status legend", level=2)
    add_para(doc, "✅ Fully Working / Implemented — complete chain verified in code; handles core errors/edges.")
    add_para(doc, "⚠️ Partially Working / Partial — happy path only, missing validation, weak auth scope, or FE/BE mismatch.")
    add_para(doc, "❌ Not Working / Missing — broken, dummy/hardcoded, missing handler, unmounted route, or UI-only.")
    add_para(doc, "⚠️ Needs manual testing — logic present but requires live app/SMTP/device proof.")

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Completion ~{completion}% | OK {fully} | Partial/Manual {partial} | Broken {broken} | n={total}")
    print(f"Roles Yes/Partial/No: {roles_yes}/{roles_partial}/{roles_no} | Security OK/Partial/Missing: {sec_ok}/{sec_partial}/{sec_missing}")


if __name__ == "__main__":
    build()
