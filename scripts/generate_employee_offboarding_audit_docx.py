"""
Employee Offboarding discovery audit — documentation only (as-built).
Output: Employee_Offboarding_Audit.docx at repo root.
"""
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_PATH = Path(__file__).resolve().parents[1] / "Employee_Offboarding_Audit.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for p in h.runs:
        p.font.size = Pt(14 if level == 1 else 12)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_table(doc: Document, headers, rows) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = str(val) if val is not None else ""


def main() -> None:
    doc = Document()
    title = doc.add_paragraph("PeopleFlow / EMS — Employee Offboarding Audit (As-Built)")
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(16)
    add_para(
        doc,
        f"Generated: {date.today().isoformat()}. Multi-tenant via organizationId. "
        "Confirmed: offboarding Mongoose schemas live in server/models/onboarding.model.js alongside onboarding; "
        "shared lifecycle infrastructure in lifecycle.model.js, lifecycleGap.model.js, lifecycleConfigService, "
        "lifecycleWorkflowRuntimeService. Discovery only — no new features proposed.",
    )

    # ── 1. Data Model ──
    add_heading(doc, "1. Data Model", level=1)
    add_para(
        doc,
        "Relationship to onboarding: Same file (onboarding.model.js) exports OnboardingWorkflow and OffboardingWorkflow as "
        "separate schemas — not a shared base class. Patterns are parallel (embedded tasks, notes, activityFeed, progress) "
        "but field names differ (onboarding uses checklist/documents/trainingModules; offboarding uses clearanceChecklist, "
        "knowledgeTransfer, accessRevocation, finalSettlement, experienceLetter). OffboardingWorkflow.assets can reference "
        "onboarding via onboardingWorkflowId / onboardingAssetId (synced from OnboardingWorkflow in offboardingService.synchronizeOffboardingAssets).",
    )

    add_heading(doc, "1.1 OffboardingWorkflow", level=2)
    add_table(
        doc,
        ["Field", "Type", "Purpose / enums"],
        [
            ("organizationId", "ObjectId", "Tenant (required, indexed)"),
            ("employeeId", "ObjectId", "Separating employee (required, indexed)"),
            ("lastWorkingDate", "Date", "Required — anchor for due dates and notice calc"),
            ("resignationDate", "Date", "Optional separation start for notice period"),
            ("noticePeriod", "Object", "startDate, policyDays, servedDays, remainingDays, buyoutApproved, buyoutDays, leaveAdjustmentDays; status: Not Applicable | Active | Completed | Waived | Buyout"),
            ("startDate, expectedCompletionDate, actualCompletionDate", "Date", "Workflow timeline"),
            ("reason", "enum", "Resignation | Termination | Retirement | Contract End | Other"),
            ("reasonDetails", "String", "Free text"),
            ("status", "enum", "Draft | Scheduled | Initiated | In Progress | Completed | Paused | Cancelled | Overdue | On Hold"),
            ("pausedFromStatus", "String", "Prior status when paused"),
            ("assignedHR, assignedManager", "ObjectId", "Workflow owners"),
            ("tasks[]", "offboardingTaskSchema", "Role-assigned tasks"),
            ("clearanceChecklist[]", "Subdoc", "Department clearance items (category HR/IT/Finance/Manager/Employee/…; assignedRole employee|hr|manager|it|payroll|admin|other; completed boolean)"),
            ("exitInterview", "Object", "Legacy embedded interview + sessions[] → ExitInterviewSession; status: Not Scheduled | Scheduled | In Progress | Completed | Waived | Overdue"),
            ("assets[]", "Subdoc", "Return tracking; links onboarding asset; status Pending | Returned | Damaged | Lost | Waived"),
            ("accessRevocation[]", "Subdoc", "system, accessType, status Pending | Completed | Waived"),
            ("knowledgeTransfer[]", "Subdoc", "topic, transferredBy/To, status Pending | In Progress | Completed | Waived | Overdue"),
            ("finalSettlement", "Object", "F&F amounts + status Pending | In Review | Calculated | Pending Approval | Approved | Paid | Cleared | Waived; approvalInstanceId, pdfUrl"),
            ("experienceLetter", "Object", "canonical doc ref; status Pending | Issued | Waived"),
            ("documents[]", "Subdoc", "Legacy/minimal; documentSummary aggregate"),
            ("progress, progressBreakdown", "Number + object", "Overall %; breakdown: tasks, checklist, assets, knowledgeTransfer, exitInterview (legacy 100), payrollClearance, accessRevocation, byRole"),
            ("activityFeed, notes, notificationHistory", "Arrays", "Audit + HR notes (visibility enum on notes)"),
            ("initiatedBy", "ObjectId", "Who started workflow"),
        ],
    )
    add_para(
        doc,
        "Note: createEmployeeOffboardingWorkflow sets offboardingTemplateId on the document, but OffboardingWorkflow schema "
        "does not declare templateId/offboardingTemplateId — value may not persist under strict Mongoose. "
        "Template clearanceRules[] exist on OffboardingTemplate but are not copied onto workflow instances; runtime clearance is clearanceChecklist only.",
    )

    add_heading(doc, "1.2 offboardingTask (embedded)", level=2)
    add_table(
        doc,
        ["Field", "Enums / notes"],
        [
            ("category", "IT | HR | Finance | Access | Equipment | Knowledge Transfer | Documentation | Other"),
            ("assigneeRole", "employee | hr | manager | it | payroll | admin | other"),
            ("status", "Pending | In Progress | Waiting | Blocked | Completed | Cancelled | Overdue"),
            ("priority", "Low | Medium | High | Critical"),
            ("mandatory", "Boolean — drives progress %"),
            ("dueDate, startedDate, completedDate", "Scheduling"),
        ],
    )

    add_heading(doc, "1.3 OffboardingTemplate", level=2)
    add_para(
        doc,
        "organizationId, name, expectedDays, template meta (Draft/Published/Archived), stages[] (lifecycleStageSchema), "
        "tasks[], clearanceChecklist[], accessRevocation[], knowledgeTransfer[], assetRules[], clearanceRules[] (HR/IT/Finance/Admin departments), "
        "documents[], exitInterview { enabled, mandatory }, finalSettlement { enabled, mandatory }, exitTypes[].",
    )

    add_heading(doc, "1.4 Related models", level=2)
    add_table(
        doc,
        ["Model", "Role in offboarding"],
        [
            ("LifecycleSettings", "modules.offboarding, noticePeriodPolicy, fnfPolicy, exitInterviewPolicy, enabledFeatures"),
            ("LifecycleFormDefinition / LifecycleFormResponse", "Exit/offboarding forms (category offboarding, exit_interview)"),
            ("ExitInterviewSession", "lifecycle.model.js — session-based exit interviews linked from workflow.exitInterview.sessions"),
            ("BackgroundVerificationCase", "Primarily onboarding; not core offboarding path"),
            ("Document", "Experience letter upload; employee document library; documentSummary sync"),
            ("lifecycleGap letter templates", "Generated letters: experience_letter, relieving_letter, fnf-related types"),
        ],
    )

    # ── 2a Employee APIs ──
    add_heading(doc, "2. Backend APIs", level=1)
    add_heading(doc, "2a. Employee-facing APIs", level=2)
    add_para(
        doc,
        "Mounts: /api/v1/both/onboarding/offboarding/* (isAuthenticatedBoth — admin or employee JWT) and "
        "/api/v1/both/lifecycle/employee/* (isAuthenticated — requires Employee profile). "
        "No dedicated employee API to submit resignation or create a workflow — separation is initiated by HR (admin route).",
    )
    add_table(
        doc,
        ["Method", "Path", "Purpose", "Request body", "Response", "Guard"],
        [
            ("GET", "/both/onboarding/offboarding/my-workflow", "Own active offboarding workflow", "—", "{ success, data: workflow (sanitized) }", "isAuthenticatedBoth"),
            ("GET", "/both/onboarding/offboarding/dashboard", "Self/team/manager stats + workflows", "?view=team", "{ workflows, stats, role }", "isAuthenticatedBoth"),
            ("GET", "/both/onboarding/offboarding/workflow/:id", "Single workflow if authorized", "—", "Populated OffboardingWorkflow", "isAuthenticatedBoth; canAccessOffboardingWorkflow"),
            ("PUT", "/both/onboarding/offboarding/workflow/:id/tasks/:taskId", "Update assigned task", "{ status, completionNotes }", "Workflow", "Assignee or employee-role task owner"),
            ("PUT", ".../checklist/:itemId", "Toggle clearance checklist item", "{ completed, notes }", "Workflow", "canUpdateOffboardingRoleItem (employee can complete employee-role items)"),
            ("PUT", ".../knowledge-transfer/:transferId", "Update KT item", "{ status, notes }", "Workflow", "Employee if transferredBy/To or HR"),
            ("PUT", ".../exit-interview", "Answer questions / submit", "{ questions[], feedback, rating, submit:true }", "Workflow", "Employee can answer; HR schedules status"),
            ("PUT", ".../assets/:assetId", "Asset return", "{ status, condition, notes }", "403 for employees", "HR or IT only on both/onboarding route"),
            ("PUT", ".../access-revocation/:accessId", "Mark access revoked", "{ status, notes }", "403 for employees", "HR or IT only on both/onboarding route"),
            ("PUT", ".../payroll-clearance", "F&F clearance flag", "{ status, notes }", "403 for employees", "HR or payroll only"),
            ("GET", "/both/lifecycle/employee/offboarding", "Alternate fetch", "—", "Enriched workflow", "requireLifecycleModule(offboarding)"),
            ("PUT", "/both/lifecycle/employee/offboarding/:wfId/tasks/:taskId", "Employee task only", "{ status, notes }", "Workflow", "employeeId scoped"),
            ("PUT", ".../checklist/:itemId", "Employee checklist", "{ completed }", "Workflow", "BUG: service uses workflow.checklist — schema field is clearanceChecklist"),
            ("PUT", ".../assets/:assetId", "Employee asset update", "{ status, notes }", "Workflow", "Conflicts with both/onboarding IT-only rule"),
            ("PUT", ".../knowledge-transfer/:id", "Employee KT", "{ status, notes }", "Workflow", "Employee-owned workflow"),
            ("PUT", ".../exit-interview", "Employee exit interview", "{ responses, feedback, rating }", "Workflow", "Uses exitInterview.responses — schema uses questions[]"),
            ("PUT", ".../access-revocation/:id", "Employee access update", "{ status, notes }", "Workflow", "No role guard — employee can set Completed (inconsistent with both route)"),
            ("POST", ".../notes", "Add note", "{ note, visibility? }", "Workflow", "Employee note uses content/isPrivate — schema uses note/visibility"),
            ("POST", "/both/lifecycle/employee/workflows/:wfId/tasks/:taskId/complete", "Complete employee task", "{ notes }", "Workflow", "Supports offboarding tasks"),
            ("GET", "/both/lifecycle/exit-interviews/mine", "List exit interview sessions", "—", "Sessions", "isAuthenticated"),
            ("GET/POST", "/both/lifecycle/exit-interviews/:id", "Session detail / draft / submit", "form responses", "Session", "Parallel to embedded workflow.exitInterview"),
        ],
    )

    add_heading(doc, "2b. Admin / HR-facing APIs", level=2)
    add_table(
        doc,
        ["Method", "Path", "Purpose", "Guard"],
        [
            ("GET/POST", "/admin/onboarding/offboarding", "List / create workflow", "isAdminOrHROrExecutiveAuthenticated"),
            ("GET/PUT/DELETE", "/admin/onboarding/offboarding/:id", "CRUD workflow settings", "HR/Admin"),
            ("POST", "/admin/onboarding/offboarding/:id/actions/:action", "publish | pause | resume | cancel", "HR/Admin"),
            ("GET", "/admin/onboarding/offboarding/dashboard", "Org dashboard", "HR/Admin"),
            ("GET", "/admin/onboarding/offboarding/analytics", "Analytics aggregates", "HR/Admin"),
            ("PUT", "/admin/onboarding/offboarding/:id/tasks|checklist|assets|knowledge-transfer|exit-interview|payroll-clearance|access-revocation/:…", "Same handlers as both route with HR privileges", "HR/Admin/role assignees"),
            ("POST", "/admin/onboarding/offboarding/:id/notes", "Internal notes", "HR/Admin"),
            ("POST", "/admin/onboarding/offboarding/:id/experience-letter", "Upload experience letter (multipart)", "HR/Admin"),
            ("GET/PUT", "/admin/lifecycle/templates/offboarding/*", "Template CRUD, stages, publish, clone", "HR/Admin"),
            ("GET/PATCH/POST", "/admin/lifecycle/offboarding/:workflowId/stages/*", "Stage runtime (same pattern as onboarding — workflow.stages not in schema)", "HR/Admin"),
            ("POST", "/admin/lifecycle/offboarding/:workflowId/notice-period/buyout", "Approve notice buyout", "noticePeriod feature"),
            ("PATCH", ".../notice-period/leave-adjustment", "Adjust notice via leave days", "noticePeriod"),
            ("POST", ".../fnf/calculate", "Run F&F calculation", "fnfSettlement"),
            ("POST", ".../fnf/submit-approval", "Approval workflow for F&F", "fnfSettlement"),
            ("POST", ".../fnf/pdf", "Generate F&F PDF", "fnfSettlement"),
            ("POST", ".../fnf/payroll-bridge", "Payroll run bridge", "fnfSettlement"),
            ("POST", ".../access/auto-revoke", "Bulk access revoke", "accessRevocation"),
            ("GET", ".../letters/generated", "Auto-generated exit letters", "letter features"),
            ("GET", "/admin/lifecycle/dashboard/notice-period", "Notice period KPI dashboard", "HR/Admin"),
            ("GET", "/admin/lifecycle/dashboard/fnf", "F&F pending dashboard", "HR/Admin"),
            ("GET/POST/PATCH", "/admin/lifecycle/exit-interviews/*", "Templates, sessions, schedule, analytics", "exitInterview feature"),
            ("POST", "/admin/lifecycle/exit-interviews/provision/:workflowId", "Provision sessions for workflow", "HR/Admin"),
            ("GET/POST", "/admin/lifecycle/alumni/*", "Alumni list, blacklist, rehire", "alumniManagement"),
            ("GET", "/both/lifecycle/manager/offboarding", "Manager team offboarding view", "Manager + offboarding module"),
        ],
    )

    add_heading(doc, "2c. Cross-check gaps", level=2)
    add_table(
        doc,
        ["Gap type", "Detail"],
        [
            ("HR-only, no employee equivalent", "Create/initiate offboarding, delete workflow, analytics, F&F calculate/approve/PDF, notice buyout/leave adjustment, asset return (both route), access revocation updates, experience letter upload, auto-revoke, alumni, template/stage admin"),
            ("Employee read via workflow GET", "Employee sees clearanceChecklist, noticePeriod, finalSettlement status (read-only in UI), assets list, access items — no separate clearance-status API"),
            ("Employee partial write", "Tasks, employee-role checklist, knowledge transfer, exit interview answers; NOT resignation submission API"),
            ("Employee endpoint without HR mirror", "lifecycle/employee/offboarding asset & access PUT — broader than both/onboarding (HR should own IT actions)"),
            ("HR review of employee input", "Exit interview: both PUT on workflow + lifecycle exit-interview sessions (HR submit/close on admin routes); KT and checklist visible on shared workflow GET"),
            ("Duplicate stacks", "both/onboarding/offboarding vs both/lifecycle/employee/offboarding — same domain, different validation and field bugs on lifecycle path"),
            ("Pattern reuse", "Tasks/checklist/notes/activityFeed mirror onboarding; offboarding adds clearanceChecklist, KT, accessRevocation, finalSettlement, experienceLetter, noticePeriod — no onboarding-style documents[] requirement array on workflow"),
        ],
    )

    # ── 3. Workflow logic ──
    add_heading(doc, "3. Offboarding Steps / Workflow Logic", level=1)
    add_para(
        doc,
        "Default template stages (offboardingDefaults.DEFAULT_OFFBOARDING_STAGES): initiation → handover → clearance → settlement → letters. "
        "createEmployeeOffboardingWorkflow seeds tasks from template/DEFAULT_OFFBOARDING_TASKS, clearanceChecklist, accessRevocation, knowledgeTransfer, exitInterview (if feature+template enabled), experienceLetter Pending, finalSettlement Pending (mandatory if fnf feature).",
    )
    add_table(
        doc,
        ["Step (backend-defined)", "Owner roles", "Mandatory default?"],
        [
            ("Submit resignation / handover info (task)", "employee", "Yes"),
            ("Knowledge transfer (task + KT records)", "employee + manager approve task", "Yes"),
            ("Collect assets (task auto-completes when assets returned)", "it", "Yes"),
            ("Revoke access (task auto-completes when access items done)", "it", "Yes"),
            ("Clearance checklist (HR/IT/Finance/Manager/Employee items)", "per assignedRole", "Most items mandatory in DEFAULT_OFFBOARDING_CHECKLIST"),
            ("Exit interview", "employee + HR schedule", "Template-dependent; waived if feature off"),
            ("Experience letter issued", "hr", "Yes (task)"),
            ("F&F / payroll clearance", "payroll/finance", "Yes if fnfSettlement enabled"),
            ("Notice period tracking", "system calc from resignationDate + lastWorkingDate", "Feature noticePeriod"),
        ],
    )
    add_para(
        doc,
        "Department clearance: Implemented as clearanceChecklist rows with category (HR, IT, Finance, Manager, Employee) and assignedRole — not separate clearanceRules documents on the workflow. "
        "Template clearanceRules (HR/IT/Finance/Admin departments) are configuration-only unless expanded into checklist at template build time.",
    )
    add_para(
        doc,
        "Gating: Progress is weighted (tasks 25%, checklist 25%, assets 15%, KT 15%, payrollClearance 10%, accessRevocation 10%) — feature gates may zero dimensions. "
        "Workflow status → Completed when progress === 100 (no explicit 'all clearances before F&F' server gate beyond progress math). "
        "payrollClearance progress counts 100 only when finalSettlement.status in Paid/Cleared/Waived — Calculated/Approved do not complete that slice. "
        "synchronizeOffboardingDerivedTasks links task completion to assets, KT, access, experience letter, F&F status.",
    )
    add_para(
        doc,
        "Notice period: noticePeriodService.applyNoticePeriodToWorkflow uses LifecycleSettings.noticePeriodPolicy (default/probation/grade/department/employment type days). "
        "calculateNoticePeriod computes servedDays, remainingDays, status Active|Completed|Buyout. Buyout and leave adjustment via HR-only admin lifecycle routes.",
    )
    add_para(
        doc,
        "Onboarding link: synchronizeOffboardingAssets pulls Assigned assets from employee's OnboardingWorkflow into offboarding.assets for return tracking.",
    )

    # ── 4. Frontend ──
    add_heading(doc, "4. Existing Frontend", level=1)
    add_heading(doc, "4a. Employee side", level=2)
    add_table(
        doc,
        ["Surface", "Route", "Built?"],
        [
            ("OffboardingHub (employee mode)", "/onboarding?lifecycle=offboarding", "Built — my-workflow, tasks, clearance checklist toggle (employee items), KT, notice period read-only, F&F status read-only, exit interview panel"),
            ("Resignation submission form", "—", "Missing — HR initiates via create modal (admin only)"),
            ("Dedicated employee offboarding page", "—", "No separate route; shared EmployeeLifecycle shell"),
            ("Exit interview (employee)", "Embedded ExitInterviewWorkflowPanel + /both/lifecycle/exit-interviews/mine", "Partial — HR panel uses admin sessions API; employees may use session APIs if linked"),
        ],
    )

    add_heading(doc, "4b. Admin/HR side", level=2)
    add_table(
        doc,
        ["Surface", "Built?"],
        [
            ("OffboardingHub HR list + initiate modal", "Built — admin/offboarding dashboard + analytics"),
            ("Clearance / assets / access sections", "Built — HR/IT actions in hub"),
            ("Notice period buyout / leave adjustment", "Built — admin lifecycle endpoints in hub"),
            ("F&F calculate, approval, PDF, payroll bridge", "Built in hub (HR only)"),
            ("Experience letter upload", "Built"),
            ("ExitInterviewWorkflowPanel, ExitInterviewHub", "Built (admin lifecycle)"),
            ("OffboardingLettersPanel, OffboardingAlumniPanel", "Built"),
            ("LifecycleReportsHub offboarding tab, EnterpriseCommandCenter KPIs", "Built"),
            ("WorkflowBuilder / LifecycleConfigurationPanel offboarding templates", "Built"),
        ],
    )

    add_heading(doc, "Design system", level=2)
    add_para(
        doc,
        "Employee offboarding uses same shell as onboarding: Tailwind, theme-primary (#20C997 in index.css, not #1CCA9B), theme-secondary blue, gray cards, Fa icons. "
        "Manrope + corporateModern (.cm-hrms) applies when HR/admin views lifecycle with enterprise chrome — same inconsistency as onboarding audit.",
    )

    # ── 5. Gaps ──
    add_heading(doc, "5. Gaps & Observations", level=1)
    add_table(
        doc,
        ["Category", "Observation"],
        [
            ("Model, no API", "Template clearanceRules[]; workflow template id; workflow.stages runtime; exitInterview.sessions provisioning details for employee-only API"),
            ("Model, no employee API", "finalSettlement calculation fields — employee read-only via workflow GET, no dedicated F&F statement download for employee"),
            ("API, no UI", "lifecycle/employee/offboarding/* largely unused by OffboardingHub (hub uses BOTH_OFFBOARDING_ENDPOINT); manager/lifecycle/manager/offboarding"),
            ("API duplication", "Parallel both/offboarding and lifecycle/employee/offboarding with inconsistent authorization"),
            ("Schema/service bugs", "lifecycleEmployeeService: checklist vs clearanceChecklist; payrollClearance vs finalSettlement; notes content vs note; exitInterview.responses vs questions[]; populate firstName/lastName on Employee"),
            ("Controller vs schema", "updateOffboardingPayrollClearance allows subset of statuses; UI exposes full finalSettlement enum"),
            ("createOffboardingWorkflow", "Sets offboardingTemplateId not on OffboardingWorkflow schema"),
            ("No employee resignation", "Feature flag resignation exists but no employee POST resignation — HR creates workflow with reason Resignation"),
            ("Asset return", "Employee lifecycle PUT can update assets; both/onboarding blocks employees — contradictory"),
        ],
    )

    add_heading(doc, "Open Questions", level=1)
    for q in [
        "Should employees self-initiate resignation (API + UI) or remain HR-initiated only?",
        "Canonical API stack: both/onboarding/offboarding vs both/lifecycle/employee/offboarding — merge plan?",
        "Should employees download F&F PDF / view calculation breakdown, or HR-only?",
        "Are template clearanceRules meant to materialize into clearanceChecklist at workflow creation?",
        "Should access revocation and asset return be employee acknowledgment only vs IT-only updates?",
        "Exit interview: single embedded workflow.exitInterview vs ExitInterviewSession — which is source of truth for UI?",
        "When should employee account deactivate / access revoke relative to lastWorkingDate (auto-revoke policy)?",
        "Confirm brand alignment (#1CCA9B, #FF986A, Manrope) for employee offboarding shell.",
    ]:
        doc.add_paragraph(q, style="List Bullet")

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
