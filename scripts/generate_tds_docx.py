"""Generate June 2026 Payroll TDS Explanation Word document (non-tech + tech)."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold)


def add_para(doc, text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, size=10)
        shade_cell(cell, "1F4E79")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            set_cell_text(cell, val, size=10)
            if r_idx % 2 == 1:
                shade_cell(cell, "F2F2F2")
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Payroll TDS Explanation")
    set_run_font(r, size=22, bold=True, color=(31, 78, 121))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("June 2026 Salary — Income Tax Deduction Walkthrough")
    set_run_font(r, size=14, color=(89, 89, 89))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "TDS deducted: ₹4,44,599  |  First payroll of FY 2026-27  |  "
        "Old tax regime  |  Independent TDS mode"
    )
    set_run_font(r, size=10, color=(89, 89, 89))

    add_para(doc, "This document has two parts:")
    add_para(doc, "Part A — Non-technical: plain language for business / HR / employees.")
    add_para(doc, "Part B — Technical: formulas, tax slabs, and system logic for engineering / payroll ops.")

    # ========== PART A ==========
    doc.add_page_break()
    doc.add_heading("PART A — Non-Technical Explanation", level=1)
    add_para(doc, "Simple language — what happened and why.")

    doc.add_heading("1. What is this about?", level=2)
    add_para(
        doc,
        "This explains how income tax (TDS — Tax Deducted at Source) was calculated "
        "for the June 2026 salary. TDS is tax taken from salary every month so a large "
        "bill does not come only at year-end.",
    )

    doc.add_heading("2. Snapshot of this month's pay", level=2)
    add_table(
        doc,
        ["Particular", "Amount"],
        [
            ["Monthly salary package used", "₹15,00,000"],
            ["Days unpaid (absent on working days)", "4 days"],
            ["Salary actually earned this month", "₹13,00,000"],
            ["Provident Fund (employee)", "₹1,800"],
            ["Income tax (TDS)", "₹4,44,599"],
            ["Total deductions", "₹4,46,399"],
            ["Money in hand (net salary)", "₹8,53,601"],
        ],
        [4.5, 2.0],
    )

    doc.add_heading("3. Why was salary ₹13 lakh, not ₹15 lakh?", level=2)
    add_para(
        doc,
        "June has 30 calendar days. Out of those, 4 working days were unpaid. "
        "Pay was reduced as:",
    )
    add_para(doc, "₹15,00,000 x 26 paid days / 30 days = ₹13,00,000", bold=True)

    doc.add_heading("4. How was the salary split?", level=2)
    add_table(
        doc,
        ["Component", "Amount", "In simple words"],
        [
            ["Basic", "₹6,50,000", "Half of earned salary (50% of gross)"],
            ["House Rent Allowance (HRA)", "₹2,60,000", "40% of Basic (not 40% of total CTC)"],
            ["Special Allowance", "₹3,90,000", "Whatever is left after Basic + HRA"],
            ["Total", "₹13,00,000", "Earned taxable salary"],
        ],
        [2.5, 1.5, 3.0],
    )
    add_para(
        doc,
        'Note: HRA looks like 20% of total salary only because 40% of (50% Basic) = 20% of CTC. '
        'The rule used is still "HRA = 40% of Basic".',
    )

    doc.add_heading("5. How does monthly TDS work? (simple idea)", level=2)
    add_para(doc, "Step 1 — Estimate total income for the full financial year (April–March).")
    add_para(doc, "Step 2 — Calculate full-year tax on that estimate (old tax regime).")
    add_para(doc, "Step 3 — Take this month's share of that yearly tax.")
    add_para(doc, "For this company's setting (Independent mode):")
    add_para(
        doc,
        "This month's tax = Yearly tax x (This month's salary / Estimated yearly salary)",
        bold=True,
    )

    doc.add_heading("6. Yearly income estimate used", level=2)
    add_table(
        doc,
        ["Period", "Amount", "Meaning"],
        [
            ["April–May", "₹0", "No earlier salary slips were generated"],
            ["June (this month)", "₹13,00,000", "Actual earned salary"],
            ["July–March (9 months)", "₹1,35,00,000", "Assumed full ₹15 lakh each month"],
            ["Estimated yearly income", "₹1,48,00,000", "Total used for tax calculation"],
        ],
        [2.5, 1.8, 3.2],
    )

    doc.add_heading("7. Yearly tax (simple summary)", level=2)
    add_table(
        doc,
        ["Step", "Amount"],
        [
            ["Estimated yearly income", "₹1,48,00,000"],
            ["Less: standard deduction", "₹50,000"],
            ["Less: PF benefit under Section 80C", "₹18,000"],
            ["Income on which tax is calculated", "₹1,47,32,000"],
            ["Tax after slabs + surcharge + cess", "₹50,61,592"],
        ],
        [4.5, 2.0],
    )
    add_para(
        doc,
        "No extra tax savings were applied for rent (HRA claim), medical insurance (80D), "
        "home loan, etc., because no tax declaration was available for this financial year.",
    )

    doc.add_heading("8. This month's TDS", level=2)
    add_para(
        doc,
        "₹50,61,592 x ₹13,00,000 / ₹1,48,00,000 = ₹4,44,599",
        bold=True,
    )

    doc.add_heading("9. Important note in plain English", level=2)
    add_para(
        doc,
        "Another common way is to simply do yearly tax / 12, which would be ₹4,21,799.",
    )
    add_para(
        doc,
        'Your system used the "share of this month\'s salary" method, which for this slip '
        "is ₹22,800 higher.",
    )
    add_table(
        doc,
        ["Method", "June TDS", "Difference"],
        [
            ["Independent (used by system)", "₹4,44,599", "—"],
            ["Flat yearly tax / 12", "₹4,21,799", "₹22,800 less"],
        ],
        [3.5, 1.8, 1.8],
    )
    add_para(
        doc,
        "Both methods are only ways to collect tax monthly. The final correct tax is "
        "settled for the full year (Form 16 / annual computation).",
    )

    doc.add_heading("10. Payslip landing", level=2)
    add_table(
        doc,
        ["Payslip field", "Amount"],
        [
            ["Total earnings", "₹13,00,000"],
            ["PF (employee)", "₹1,800"],
            ["Income tax (TDS)", "₹4,44,599"],
            ["Total deductions", "₹4,46,399"],
            ["In-hand salary", "₹8,53,601"],
        ],
        [4.5, 2.0],
    )

    # ========== PART B ==========
    doc.add_page_break()
    doc.add_heading("PART B — Technical Explanation", level=1)
    add_para(doc, "System logic, formulas, regime rules, and implementation references.")

    doc.add_heading("1. System configuration inferred for this run", level=2)
    add_table(
        doc,
        ["Setting", "Value"],
        [
            ["Tax regime", "Old"],
            ["Org TDS mode", "independent"],
            ["Effective CTC for month", "₹15,00,000"],
            ["YTD payrolls in FY", "None (first payroll)"],
            ["Tax declaration (FY 2026-27)", "Not applied / not present"],
            [
                "Code path",
                "computeRegimeIncomeTaxCumulative → independent monthly cut",
            ],
        ],
        [3.0, 4.0],
    )

    doc.add_heading("2. Month salary & proration", level=2)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["Effective monthly CTC", "₹15,00,000"],
            ["Calendar days (June)", "30"],
            ["Unpaid working days", "4 (23 working − 19 present)"],
            ["Paid payroll days", "26 (30 − 4)"],
            ["Earned taxable gross", "round(15,00,000 x 26 / 30) = ₹13,00,000"],
        ],
        [3.0, 4.0],
    )
    add_para(doc, "Earnings structure (salary structure rules):")
    add_para(doc, "• Basic = 50% of gross → ₹6,50,000")
    add_para(doc, "• HRA = 40% of basic → ₹2,60,000")
    add_para(
        doc,
        "• Special = residual (gross − basic − HRA − transport − medical) → ₹3,90,000",
    )
    add_para(doc, "Other deductions: PF employee ₹1,800; PT/ESI ₹0")

    doc.add_heading("3. Projected annual income (FY 2026-27)", level=2)
    add_para(
        doc,
        "ProjectedAnnualGross = YTD taxable (0) + Current month taxable (13,00,000) "
        "+ Remaining months x CTC (9 x 15,00,000) = 1,48,00,000",
    )
    add_table(
        doc,
        ["Part", "Formula", "Amount"],
        [
            ["YTD (Apr–May)", "No prior payrolls", "₹0"],
            ["Current month (June)", "Earned taxable", "₹13,00,000"],
            ["Remaining (Jul–Mar)", "9 x ₹15,00,000", "₹1,35,00,000"],
            ["Projected annual gross", "", "₹1,48,00,000"],
            ["Projected annual PF", "1,800 + (9 x 1,800)", "₹18,000"],
            ["Projected annual PT", "", "₹0"],
        ],
        [2.5, 2.5, 1.8],
    )

    doc.add_heading("4. Old regime — deductions & reliefs", level=2)
    add_table(
        doc,
        ["Relief / deduction", "Eligible?", "Amount", "Why"],
        [
            ["Standard deduction", "Yes", "₹50,000", "Old regime fixed"],
            ["Section 80C (declared)", "No", "₹0", "No FY-matched tax declaration"],
            [
                "Section 80C — Employee PF",
                "Yes",
                "₹18,000",
                "Auto from projected PF (cap ₹1.5L)",
            ],
            ["Total 80C", "", "₹18,000", "min(1,50,000, 18,000)"],
            ["Section 80D", "No", "₹0", "No declaration"],
            ["Section 80E", "No", "₹0", "No declaration"],
            ["Section 24b", "No", "₹0", "No declaration"],
            ["HRA exemption u/s 10(13A)", "No", "₹0", "Monthly rent = 0"],
            ["Professional tax", "No", "₹0", "Not deducted in payroll"],
            ["Total deductions from gross", "", "₹68,000", "50,000 + 18,000"],
        ],
        [2.4, 1.0, 1.2, 2.4],
    )
    add_table(
        doc,
        ["Step", "Amount"],
        [
            ["Projected annual gross", "₹1,48,00,000"],
            ["(−) Total deductions", "₹68,000"],
            ["Taxable income", "₹1,47,32,000"],
        ],
        [4.5, 2.0],
    )

    doc.add_heading("5. Slab tax (Old regime)", level=2)
    add_table(
        doc,
        ["Slab", "Rate", "Income in slab", "Tax"],
        [
            ["₹0 – ₹2,50,000", "0%", "₹2,50,000", "₹0"],
            ["₹2,50,000 – ₹5,00,000", "5%", "₹2,50,000", "₹12,500"],
            ["₹5,00,000 – ₹10,00,000", "20%", "₹5,00,000", "₹1,00,000"],
            ["Above ₹10,00,000", "30%", "₹1,37,32,000", "₹41,19,600"],
            ["Tax before rebate", "", "", "₹42,32,100"],
        ],
        [2.5, 1.0, 2.0, 1.5],
    )

    doc.add_heading("6. Rebate 87A / surcharge / cess", level=2)
    add_table(
        doc,
        ["Item", "Applied?", "Amount", "Rule"],
        [
            ["Rebate u/s 87A", "No", "₹0", "Only if taxable ≤ ₹5,00,000"],
            ["Tax after rebate", "", "₹42,32,100", ""],
            ["Surcharge band", "Yes", "15%", "Taxable between ₹1 Cr and ₹2 Cr"],
            ["Raw surcharge", "", "₹6,34,815", "42,32,100 x 15%"],
            ["Marginal relief", "Not needed", "₹0", "Raw surcharge within relief cap"],
            ["Surcharge charged", "", "₹6,34,815", ""],
            ["Tax + surcharge", "", "₹48,66,915", ""],
            ["Health & education cess", "4%", "₹1,94,676.60", "48,66,915 x 4%"],
            ["Annual tax liability", "", "₹50,61,592", "Rounded"],
        ],
        [2.3, 1.2, 1.5, 2.5],
    )

    doc.add_heading("7. Monthly TDS cut (Independent mode)", level=2)
    add_para(doc, "Formula:")
    add_para(
        doc,
        "monthlyTDS = round( annualTax x currentMonthGross / projectedAnnualGross )",
        bold=True,
    )
    add_para(
        doc,
        "= round( 50,61,592 x 13,00,000 / 1,48,00,000 ) = ₹4,44,599",
        bold=True,
    )
    add_para(
        doc,
        "Past TDS is not subtracted in independent mode (and YTD TDS is ₹0 anyway).",
    )

    doc.add_heading("8. Comparison vs annual / 12", level=2)
    add_table(
        doc,
        ["Method", "Formula", "Amount"],
        [
            ["Independent (used)", "50,61,592 x 13,00,000 / 1,48,00,000", "₹4,44,599"],
            ["Flat monthly", "50,61,592 / 12", "₹4,21,799"],
            ["Difference", "Independent − flat", "₹22,800"],
        ],
        [2.2, 3.5, 1.5],
    )
    add_para(
        doc,
        "Independent mode ties TDS to this month's earned gross. Flat /12 spreads tax "
        "equally by calendar month. For this slip, missing Apr–May income in the year "
        "projection also makes June's independent share slightly larger than 1/12.",
    )

    doc.add_heading("9. Payslip landing (technical)", level=2)
    add_table(
        doc,
        ["Field", "Amount"],
        [
            ["Total earnings / CTC stored", "₹13,00,000"],
            ["PF employee", "₹1,800"],
            ["Income tax (TDS)", "₹4,44,599"],
            ["Total deductions", "₹4,46,399"],
            ["In-hand salary", "₹8,53,601"],
        ],
        [4.0, 2.0],
    )

    doc.add_heading("10. Quick formula summary", level=2)
    for line in [
        "Gross (projected FY)     ₹1,48,00,000",
        "− Std deduction             ₹50,000",
        "− 80C (PF only)             ₹18,000",
        "= Taxable                 ₹1,47,32,000",
        "→ Slab tax                 ₹42,32,100",
        "− 87A rebate                      ₹0",
        "+ Surcharge 15%             ₹6,34,815",
        "+ Cess 4%                   ₹1,94,677",
        "= Annual tax               ₹50,61,592",
    ]:
        add_para(doc, line, space_after=2)
    add_para(doc, "→ June TDS (independent)    ₹4,44,599", bold=True)

    doc.add_heading("11. Implementation references", level=2)
    add_table(
        doc,
        ["Area", "File / function"],
        [
            [
                "Annual tax",
                "server/helpers/incomeTaxCalculation.js → computeAnnualIncomeTax",
            ],
            [
                "Monthly TDS cut",
                "server/helpers/tdsCumulativeHelper.js → computeIndependentModeMonthlyTds",
            ],
            [
                "Earnings split",
                "server/helpers/salaryCalculation.js → computeSalaryStructureCoreEarnings",
            ],
            [
                "Payroll wiring",
                "server/helpers/payrollGenerationHelper.js → computeRegimeIncomeTax",
            ],
        ],
        [2.0, 5.0],
    )

    doc.add_heading("12. Document control", level=2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Period", "June 2026"],
            ["Payroll type", "Monthly / correction / final"],
            ["Audience", "Business (Part A) + Engineering (Part B)"],
            [
                "Note",
                "Figures reverse-engineered from payroll amounts and EMS TDS formulas",
            ],
        ],
        [2.0, 5.0],
    )

    out = Path(__file__).resolve().parents[1] / "June_2026_Payroll_TDS_Explanation.docx"
    doc.save(str(out))
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
