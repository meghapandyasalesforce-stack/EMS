from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt


OUT_PATH = Path(__file__).resolve().parents[1] / "Onboarding_Offboarding_Modules_Existing_Implementation_Analysis_DETAILED.docx"

F = "Fully Implemented"
P = "Partially Implemented"
M = "Not Implemented"


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(16)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.size = Pt(12 if level >= 2 else 14)


def add_table(doc: Document, headers, rows) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = str(val) if val is not None else ""


def fmt(status: str, ui: str, api: str, realtime: str, config: str) -> str:
    return f"{status}\nUI: {ui}\nAPI: {api}\nRealtime: {realtime}\nConfigure: {config}"


def main() -> None:
    doc = Document()
    add_title(doc, "EMS - Admin & HR Onboarding/Offboarding Analysis (Detailed, Pre-UI Redesign)")
    add_para(doc, f"Generated: {date.today().isoformat()}")
    add_para(
        doc,
        "Method: static tracing from existing React components to backend endpoints they call (no runtime). "
        "Realtime signals: socket.io usage in hubs (`onboarding:changed`, `offboarding:changed`). "
        "Configure integration: lifecycle feature gating via `useLifecycleFeatures` / `isFeatureEnabled` / `filterByFeature` and registry settings PATCH in `LifecycleConfigurationPanel`.",
    )

    add_heading(doc, "Step 1 - Entry points & module coverage", level=1)
    add_para(doc, "Onboarding lifecycle hub (HR/Admin): `vite_client/src/HR/Onboarding/OnboardingHub.jsx`")
    add_para(doc, "Onboarding sub-panels: `OnboardingBgvPanel`, `OnboardingFormsPanel`, `OnboardingGoalsPanel`, `OnboardingMilestonesPanel`, `DocumentHub`, `ProbationCenter`, `ItProvisioningPanel`, `LmsDashboard`")
    add_para(doc, "Offboarding lifecycle hub (HR/Admin): `vite_client/src/HR/Offboarding/OffboardingHub.jsx`")
    add_para(doc, "Offboarding sub-panels: `ExitInterviewWorkflowPanel`, `OffboardingLettersPanel`, `OffboardingAlumniPanel`")
    add_para(doc, "Admin Onboarding legacy UI: `vite_client/src/HR/Onboarding/OnboardingPage.jsx` + `OnboardingDetailPage.jsx` (legacy workflow/task/BGV UI)")
    add_para(doc, "Enterprise/Configure/Reports shell: `vite_client/src/HR/Onboarding/EmployeeLifecycle.jsx` -> `LifecycleEnterpriseHub`, `LifecycleConfigurationPanel`, `LifecycleReportsHub`")

    add_heading(doc, "Step 2 - Feature coverage matrices (UI/backend/APIs/realtime/config)", level=1)

    # Onboarding (Admin legacy vs HR lifecycle)
    onb_features = [
        "Dashboard",
        "Statistics Cards",
        "Workflow Stages",
        "Employee Details",
        "Search",
        "Filters",
        "Documents",
        "Document Hub",
        "Forms",
        "Form Builder",
        "Tasks",
        "Checklist",
        "Background Verification",
        "Buddy",
        "Goals",
        "Probation",
        "30-60-90 Plan",
        "Software Requests",
        "IT Assets",
        "Training/LMS",
        "Timeline",
        "Activity Log",
        "Notes",
        "Comments",
        "Approvals",
        "Notifications",
        "Reports",
        "Workflow Builder",
        "Offer Letter",
        "Candidate Details",
        "Configure",
        "Enterprise",
        "Progress Tracking",
        "Status Cards",
        "Buttons",
        "Tables",
        "Empty States",
        "Loading States",
    ]

    admin_onb_rows = {
        "Dashboard": fmt(P, "UI: legacy onboarding workflow cards (grid list) in OnboardingPage", "GET /v1/admin/onboarding/onboarding", "No hub socket for legacy UI (state updates after API calls).", "No lifecycle gating in legacy UI."),
        "Statistics Cards": fmt(M, "Not present in legacy OnboardingPage.", "N/A", "N/A", "N/A"),
        "Workflow Stages": fmt(P, "UI: card status + statusFilter dropdown", "GET /v1/admin/onboarding/onboarding", "No socket realtime.", "No lifecycle gating."),
        "Employee Details": fmt(P, "UI: OnboardingDetailPage header + sidebar (HR/buddy/welcome/start/end)", "GET /v1/admin/onboarding/onboarding/:id", "No socket realtime.", "No lifecycle gating."),
        "Search": fmt(F, "UI: searchTerm filter on employee name/email", "GET /v1/admin/onboarding/onboarding", "No socket realtime.", "No lifecycle gating."),
        "Filters": fmt(F, "UI: statusFilter dropdown", "GET /v1/admin/onboarding/onboarding", "No socket realtime.", "No lifecycle gating."),
        "Documents": fmt(M, "No documents section in legacy UI.", "N/A (not used in legacy pages).", "N/A", "N/A"),
        "Document Hub": fmt(M, "Not present in legacy UI.", "N/A", "N/A", "N/A"),
        "Forms": fmt(M, "No mandatory forms UI in legacy pages.", "N/A", "N/A", "N/A"),
        "Form Builder": fmt(M, "Only appears in enterprise Configure hub (FormBuilderHub).", "POST/PUT /v1/admin/lifecycle/forms (configure area)", "N/A for legacy pages.", "No configure in legacy pages."),
        "Tasks": fmt(F, "UI: tasks list + add task + Start/Complete in OnboardingPage/OnboardingDetailPage", "PUT /v1/admin/onboarding/onboarding/:id/tasks/:taskId; PUT /v1/admin/onboarding/onboarding/:id", "No socket realtime; state updated from API responses.", "No lifecycle gating."),
        "Checklist": fmt(M, "No checklist UI in legacy pages.", "N/A", "N/A", "N/A"),
        "Background Verification": fmt(P, "UI: BGV case panel in OnboardingDetailPage (initiate + HR status dropdown for checks)", "GET/POST ${ADMIN_LIFECYCLE_ENDPOINT}/bgv/hr (legacy uses admin/lifecycle/bgv); see OnboardingDetailPage calls", "No socket realtime for BGV edits (reload via fetchWorkflow).", "No lifecycle feature gating in legacy UI."),
        "Buddy": fmt(P, "UI: buddy displayed in sidebar only (no check-ins list).", "Included in workflow payload; no buddy-checkin endpoints in legacy UI.", "No socket realtime.", "No lifecycle gating."),
        "Goals": fmt(M, "No goals UI.", "N/A", "N/A", "N/A"),
        "Probation": fmt(M, "No probation UI.", "N/A", "N/A", "N/A"),
        "30-60-90 Plan": fmt(M, "No milestones UI.", "N/A", "N/A", "N/A"),
        "Software Requests": fmt(M, "No software/IT provisioning UI.", "N/A", "N/A", "N/A"),
        "IT Assets": fmt(M, "No asset inventory UI.", "N/A", "N/A", "N/A"),
        "Training/LMS": fmt(M, "No training/LMS UI.", "N/A", "N/A", "N/A"),
        "Timeline": fmt(M, "No activity timeline UI.", "N/A", "N/A", "N/A"),
        "Activity Log": fmt(M, "No activity feed.", "N/A", "N/A", "N/A"),
        "Notes": fmt(M, "No notes UI.", "N/A", "N/A", "N/A"),
        "Comments": fmt(M, "No comment-thread UI.", "N/A", "N/A", "N/A"),
        "Approvals": fmt(M, "No approval queue UI in legacy pages.", "N/A", "N/A", "N/A"),
        "Notifications": fmt(M, "No notifications feed/UI in legacy pages.", "N/A (enterprise-only configure exists).", "N/A", "N/A"),
        "Reports": fmt(M, "Reports not in legacy pages; use LifecycleReportsHub.", "GET /v1/admin/lifecycle/reports/*", "N/A", "Configure area only."),
        "Workflow Builder": fmt(M, "WorkflowBuilder UI exists only in Configure (enterprise).", "PUT/POST /v1/admin/lifecycle/templates/onboarding|offboarding", "N/A", "Configure area only."),
        "Offer Letter": fmt(M, "Offer management is separate from legacy onboarding pages.", "Offers endpoints in OfferTemplateManager/OfferLetterPanel", "N/A", "Configure area only."),
        "Candidate Details": fmt(P, "Shows employee name/email; does not include offer-candidate portal details flow.", "GET /v1/admin/onboarding/onboarding (workflow includes employeeId)", "No socket realtime.", "No gating."),
        "Configure": fmt(M, "Not present in legacy onboarding UI.", "N/A", "N/A", "N/A"),
        "Enterprise": fmt(M, "Not present in legacy onboarding UI.", "N/A", "N/A", "N/A"),
        "Progress Tracking": fmt(P, "Progress bar + completed tasks counts in list/detail.", "Derived from workflow.tasks; relies on GET workflow payload.", "No socket realtime.", "No gating."),
        "Status Cards": fmt(P, "Status badges on workflow cards.", "Derived from workflow.status.", "No socket realtime.", "No gating."),
        "Buttons": fmt(F, "Create workflow modal + add task modal + action buttons.", "POST/PUT/DELETE legacy endpoints.", "No socket realtime.", "N/A"),
        "Tables": fmt(M, "Legacy UI uses cards/grid instead of data tables.", "N/A", "N/A", "N/A"),
        "Empty States": fmt(F, "Empty messages when no workflows/tasks.", "N/A", "N/A", "N/A"),
        "Loading States": fmt(F, "Loading spinners in pages.", "N/A", "N/A", "N/A"),
    }

    hr_onb_rows = {
        "Dashboard": fmt(F, "UI: OnboardingHub list/table + KPIs (HR/Admin view)", "GET /v1/admin/onboarding/onboarding/onboarding/dashboard?view=team|all (via ADMIN_ONBOARDING_ENDPOINT in OnboardingHub)", "Realtime: socket.io `onboarding:changed` triggers fetchDashboard({soft:true}).", "Configure: hub sections gated by `useLifecycleFeatures` + `isFeatureEnabled`."),
        "Statistics Cards": fmt(F, "UI: CmStat KPIs (pending docs/verifications, offer acceptance rate, avg onboarding time)", "Same dashboard payload (dashboard.stats).", "Realtime: socket refresh.", "Configure gating for which stats exist (via feature keys)."),
        "Workflow Stages": fmt(F, "UI: workflow status chips + status badges", "Same dashboard.workflows payload.", "Realtime: socket refresh.", "Configure: WORKFLOW_STATUSES and task/checklist filtering via filterByFeature."),
        "Employee Details": fmt(F, "UI: selected workflow header includes HR/manager/buddy + welcome message", "Same selected workflow from dashboard payload.", "Realtime: socket refresh.", "Configure: conditional rendering by feature keys."),
        "Search": fmt(F, "UI: CmSearch filters workflows client-side", "Dashboard returns workflows list.", "Realtime: socket refresh.", "N/A"),
        "Filters": fmt(F, "UI: status filter chips", "Same dashboard payload.", "Realtime: socket refresh.", "N/A"),
        "Documents": fmt(P, "UI: Documents section shows status + navigates to Documents/DocumentHub; upload/verification managed in Documents area", "Document status via workflows payload; Documents pages use /v1/admin/documents (not lifecycle hub storage).", "Realtime: hub refresh updates status after actions; Documents area itself is not socket-driven.", "Configure: documents section gated by `isFeatureEnabled('documentCollection')`."),
        "Document Hub": fmt(F, "UI: `HR/Onboarding/DocumentHub.jsx` (library/pending/expiry/categories/reports)", "GET/POST /v1/admin/lifecycle/document-hub/* (categories/library/pending/expiry/catalog/reports, verify/bulk verify/remind/export)", "Realtime: reload on actions (not socket-driven).", "Configure: shown when relevant lifecycle doc features are enabled (see LifecycleConfigurationPanel visibleSections documents-hub)."),
        "Forms": fmt(F, "UI: `OnboardingFormsPanel` mandatory forms list + fill/submit modal", "GET /v1/admin/lifecycle/onboarding/:workflowId/forms/status; GET/POST /v1/both/lifecycle/forms/:formId (definition + submit)", "Realtime: reload after submit; no socket inside panel.", "Configure: rendered when any of digitalForms/emergencyContact/familyNominee/bankDetails features are enabled."),
        "Form Builder": fmt(F, "UI: `FormBuilderHub` definition builder + versioning", "GET/PUT/POST /v1/admin/lifecycle/forms and /v1/admin/lifecycle/forms/field-types", "Realtime: reload on save.", "Configure: lifecycle form builder is in Configure area."),
        "Tasks": fmt(F, "UI: Tasks by role with Start/Continue/Mark Complete + blocker enforcement", "PUT /v1/admin/onboarding/onboarding/:wfId/tasks/:taskId; workflow actions via POST /v1/admin/onboarding/onboarding/onboarding/:wfId/actions/:action", "Realtime: socket refresh updates workflow list; local state updated from API response.", "Configure: tasks filtered via `filterByFeature(wf.tasks)`."),
        "Checklist": fmt(F, "UI: Checklist section checkboxes + toggles", "PUT /v1/admin/onboarding/onboarding/:wfId/checklist/:itemId", "Realtime: refreshWorkflow after toggle + socket refresh.", "Configure: gating via `joiningChecklist` and/or `induction` features."),
        "Background Verification": fmt(F, "UI: `OnboardingBgvPanel` with HR dashboard + checks and HR actions", "GET/POST /v1/admin/lifecycle/bgv/hr/dashboard and /v1/admin/lifecycle/bgv/hr/:caseId/detail; POST /v1/admin/lifecycle/bgv (initiate); check endpoints under /v1/admin/lifecycle/bgv/:caseId/checks/:checkId/*/verify/request-info/documents/references/declaration", "Realtime: reload on action; HR panel also updates when hub refetches.", "Configure: rendered when `backgroundVerification` is enabled; panel can be disabled via hr/dashboard response."),
        "Buddy": fmt(P, "UI: Buddy Check-ins (status dropdown) + buddy assignment displayed", "PUT /v1/admin/onboarding/onboarding/:wfId/buddy-checkins/:checkinId", "Realtime: refreshWorkflow after update; socket refresh updates hub.", "Configure: gating via `buddyProgram` feature."),
        "Goals": fmt(F, "UI: Probation Goals & KRA list + HR assigns defaults", "GET/POST /v1/admin/lifecycle/onboarding/:wfId/goals", "Realtime: reload on assign.", "Configure: `goalAssignment` feature enables section."),
        "Probation": fmt(F, "UI: `ProbationCenter` pipeline + extend + review sign-off", "GET /v1/admin/lifecycle/probation/center; POST /v1/admin/lifecycle/onboarding/:wfId/probation/extend and /review", "Realtime: manual refresh; no socket inside center.", "Configure: section enabled by `probationCenter`/`probation`."),
        "30-60-90 Plan": fmt(F, "UI: `OnboardingMilestonesPanel` day-plan review form (role-based submit)", "POST /v1/admin/lifecycle/onboarding/:wfId/milestones/review", "Realtime: reload on submit; no socket inside panel.", "Configure: gating via `milestone30_60_90` and wf.milestones length."),
        "Software Requests": fmt(F, "UI: ItProvisioningPanel workflow-scoped software requests (system input + approve/provision)", "POST /v1/admin/lifecycle/onboarding/:wfId/software-requests; POST /approve; POST /provision; plus IT providers inventory endpoints", "Realtime: reload after approve/provision.", "Configure: gating via `softwareAccessRequest` feature."),
        "IT Assets": fmt(F, "UI: Equipment & Assets section updates asset.status for HR view", "PUT /v1/admin/onboarding/onboarding/:wfId/assets/:assetId", "Realtime: refreshWorkflow after update + socket refresh.", "Configure: gating via `itAssetProvisioning` feature."),
        "Training/LMS": fmt(F, "UI: training modules list + `LmsDashboard` for org/workflow sync + completion", "PUT /v1/admin/onboarding/onboarding/:wfId/training/:trainingId (Complete); GET /v1/admin/lifecycle/lms/dashboard and /lms/providers; POST /onboarding/:wfId/lms/assign-course; POST /onboarding/:wfId/lms/sync-certify", "Realtime: reload after completion/sync; hub socket refresh.", "Configure: gating via mandatoryTraining/trainingLms/lmsDashboard features."),
        "Timeline": fmt(P, "UI: Recent Activity list (timeline-like list) for selected workflow", "Provided by workflow.activityFeed from dashboard payload.", "Realtime: socket refresh.", "N/A"),
        "Activity Log": fmt(F, "UI: activityFeed list with actor/time rendering", "Same as above.", "Realtime: socket refresh.", "N/A"),
        "Notes": fmt(F, "UI: HR-only Internal Notes list + add textarea", "POST /v1/admin/onboarding/onboarding/:wfId/notes", "Realtime: refresh after add + socket.", "N/A"),
        "Comments": fmt(M, "No comment-thread UI; only internal notes.", "N/A", "N/A", "N/A"),
        "Approvals": fmt(P, "UI: partial approvals (offers approval + policy signature sign). No onboarding approval queue panels in hub.", "Offer approvals endpoints: /v1/admin/lifecycle/offers/:offerId/submit-approval; policy signing endpoints via hub signPolicy (uses POST /v1/both/onboarding/workflow/:wfId/signatures/:signatureId/sign)", "Realtime: local reload; approvals not socket-driven.", "Configure: approval support exists in Configure registry; hub UI does not present full approval queues."),
        "Notifications": fmt(P, "UI: notifications are configurable in enterprise hub; hubs do not render a dedicated notification feed.", "Configure: GET/PATCH /v1/admin/lifecycle/notifications/center; notification APIs exist under /v1/admin/notifications/* (separate).", "Realtime: not rendered in hub.", "Configure: `notificationCenter` feature enables notifications tab."),
        "Reports": fmt(P, "UI: hub shows KPI strip; deep reporting in LifecycleReportsHub.", "GET /v1/admin/lifecycle/reports/dashboard|catalog|scheduled|drilldown|export", "Realtime: reports are not socket-driven; manual refresh.", "Configure: reports accessible via EmployeeLifecycle query `reports=1`."),
        "Workflow Builder": fmt(F, "UI: WorkflowBuilder supports drag-and-drop stage design + publish/clone/preview", "GET/PUT/POST /v1/admin/lifecycle/templates/onboarding|offboarding and /publish", "Realtime: manual reload.", "Configure area; registry features decide which stages/features become visible."),
        "Offer Letter": fmt(F, "UI: Offer templates + offer generation + send/approval + candidate accept/decline", "Offer endpoints: GET/POST /v1/admin/lifecycle/offers/templates, GET/POST /v1/admin/lifecycle/offers/candidate/:candidateId, POST /v1/admin/lifecycle/offers/:offerId/send and /submit-approval; public portal endpoints: GET/POST /v1/public/lifecycle/offer/:token", "Realtime: none; actions reload offers.", "Configure: offer letter/approval features enabled via registry."),
        "Candidate Details": fmt(P, "UI: hub shows candidate/workflow progression; offer-candidate portal details are in OfferCandidatePortal", "Onboarding hub payload includes candidate status; public portal loads offer by token.", "No socket for portal.", "N/A"),
        "Configure": fmt(F, "UI: LifecycleConfigurationPanel toggles modules/policies/templates/approvals/escalations/reminders", "GET /v1/admin/lifecycle/registry; PATCH /v1/admin/lifecycle/settings", "Realtime: registry snapshot refresh after save.", "Configure: primary mechanism."),
        "Enterprise": fmt(F, "UI: LifecycleEnterpriseHub command-center modules (dashboards, statutory, IT, LMS, probation, notifications)", "Various /v1/admin/lifecycle/* endpoints per module; notifications config via /notifications/center.", "Realtime: none; manual/bulk actions reload.", "Configure: module tabs gated by enabled registry features."),
        "Progress Tracking": fmt(F, "UI: progress rings + progress breakdown per category", "Derived from workflow payload breakdown values.", "Realtime: socket refresh.", "N/A"),
        "Status Cards": fmt(F, "UI: CmStat KPI cards and status badges in hub tables", "Derived from dashboard payload.", "Realtime: socket refresh.", "N/A"),
        "Buttons": fmt(F, "UI: action buttons throughout hub/panels", "Various PUT/POST endpoints per action.", "After action: hub refresh; sometimes panel reload.", "Feature gated visibility."),
        "Tables": fmt(F, "UI: workflow/candidate list uses cm-table styling", "Derived from dashboard.workflows payload.", "Realtime: socket refresh.", "N/A"),
        "Empty States": fmt(F, "UI: CmEmpty + inline empty messages per module section", "N/A", "N/A", "N/A"),
        "Loading States": fmt(F, "UI: LoadingSpinner + CM loading states", "N/A", "N/A", "N/A"),
    }

    add_heading(doc, "Onboarding (Admin legacy vs HR lifecycle hub) - Detailed rows", level=2)
    onb_rows_table = []
    for feat in onb_features:
        onb_rows_table.append([feat, admin_onb_rows[feat], hr_onb_rows[feat]])
    add_table(doc, headers=["Feature", "Admin Onboarding (Legacy UI)", "HR Onboarding (Lifecycle Hub)"], rows=onb_rows_table)

    # Offboarding (Admin+HR)
    add_heading(doc, "Offboarding - Detailed rows (Admin + HR use OffboardingHub)", level=2)
    off_features = [
        "Dashboard",
        "Statistics Cards",
        "Exit Workflow",
        "Notice Period",
        "Exit Interview",
        "Clearance",
        "Knowledge Transfer",
        "Asset Return",
        "Access Revocation",
        "F&F Settlement",
        "Payroll Bridge",
        "Generated Letters",
        "Experience Letter",
        "Relieving Letter",
        "Service Certificate",
        "Alumni",
        "Rehire",
        "Blacklist",
        "Timeline",
        "Reports",
        "Forms",
        "Workflow",
        "Activity Logs",
        "Notifications",
        "Configure",
        "Enterprise",
        "Progress Tracking",
        "Status Cards",
        "Buttons",
        "Tables",
        "Empty States",
        "Loading States",
    ]

    off = {
        "Dashboard": fmt(F, "UI: OffboardingHub list + stats when no workflow selected", "GET /v1/admin/onboarding/offboarding/dashboard and /analytics", "Realtime: socket.io `offboarding:changed` triggers fetchDashboard({soft:true}).", "Configure gating via useLifecycleFeatures filterByFeature/isFeatureEnabled."),
        "Statistics Cards": fmt(F, "UI: KPI stats grid (Total Separations, Pending Assets, etc.)", "dashboard.stats (from /dashboard)", "Realtime: socket refresh.", "N/A"),
        "Exit Workflow": fmt(F, "UI: workflow list + selected workflow detail with publish/pause/resume/cancel actions", "POST /v1/admin/onboarding/offboarding/:wfId/actions/:action", "Realtime: socket refresh + local selectedWorkflow update.", "N/A"),
        "Notice Period": fmt(F, "UI: noticePeriod summary + HR buyout approval + leave adjustment", "POST /v1/admin/lifecycle/offboarding/:wfId/notice-period/buyout and PATCH /leave-adjustment; payload from /dashboard", "Realtime: refresh after lifecycle actions + socket.", "Configure gating via `noticePeriod` feature."),
        "Exit Interview": fmt(F, "UI: ExitInterviewWorkflowPanel sessions list + draft/submit/close + provision", "GET /v1/admin/lifecycle/exit-interviews/sessions?offboardingWorkflowId=:id; POST /draft|submit|close; POST /provision/:workflowId; GET /forms/:formDefinitionId", "Realtime: panel reload after actions; socket refresh for hub.", "Configure gating via `exitInterview` feature."),
        "Clearance": fmt(F, "UI: Clearance checklist checkboxes", "PUT /v1/admin/onboarding/offboarding/:wfId/checklist/:itemId", "Realtime: refreshWorkflow + socket.", "Configure gating via `clearance` or `hrClearance` features."),
        "Knowledge Transfer": fmt(F, "UI: knowledge transfer items with status dropdown", "PUT /v1/admin/onboarding/offboarding/:wfId/knowledge-transfer/:itemId", "Realtime: refreshWorkflow + socket.", "Configure gating via knowledgeTransfer/projectHandover features."),
        "Asset Return": fmt(F, "UI: Asset Return section with asset.status dropdown", "PUT /v1/admin/onboarding/offboarding/:wfId/assets/:assetId", "Realtime: refreshWorkflow + socket.", "Configure gating via `assetReturn` feature."),
        "Access Revocation": fmt(F, "UI: Access revocation list + auto-revoke button", "POST /v1/admin/lifecycle/offboarding/:wfId/access/auto-revoke; PUT /v1/admin/onboarding/offboarding/:wfId/access-revocation/:itemId", "Realtime: refreshWorkflow + socket.", "Configure gating via `accessRevocation` / `emailAccessRemoval` features."),
        "F&F Settlement": fmt(P, "UI: F&F calculator + submit-for-approval + generate PDF + create payroll run", "POST /v1/admin/lifecycle/offboarding/:wfId/fnf/calculate, /submit-approval, /pdf, /payroll-bridge; payload from /dashboard", "Realtime: lifecycle actions trigger re-fetch; socket refresh.", "Configure gating via `fnfSettlement` feature."),
        "Payroll Bridge": fmt(P, "UI: payroll bridge button exists but no detailed payroll-run status page in hub", "POST /v1/admin/lifecycle/offboarding/:wfId/fnf/payroll-bridge", "Realtime: refresh after create.", "Configure gating via fnfSettlement feature."),
        "Generated Letters": fmt(F, "UI: OffboardingLettersPanel shows auto-generated letters list and download", "GET /v1/admin/lifecycle/offboarding/:wfId/letters/generated", "Realtime: panel reload after lifecycle completion (not socket-driven).", "Configure gating via experienceLetter/relievingLetter/serviceCertificate features."),
        "Experience Letter": fmt(F, "UI: dedicated upload/replace + view for experience letter", "POST /v1/admin/onboarding/offboarding/:wfId/experience-letter; payload from /dashboard + GET in panel", "Realtime: upload triggers hub refresh.", "Configure gating via `experienceLetter` feature."),
        "Relieving Letter": fmt(P, "UI: only appears in generic generated letters list (no dedicated rel./cert section)", "Letters generated via /letters/generated", "Realtime: reload after generation.", "Configure gating via `relievingLetter` feature."),
        "Service Certificate": fmt(P, "UI: only appears in generic generated letters list (no dedicated section)", "Letters generated via /letters/generated", "Realtime: reload after generation.", "Configure gating via `serviceCertificate` feature."),
        "Alumni": fmt(F, "UI: OffboardingAlumniPanel shows alumni records + mini KPI cards (HR-only)", "GET /v1/admin/lifecycle/alumni; POST /alumni/:id/rehire; POST /alumni/:id/blacklist", "Realtime: panel reload after actions; no socket feed.", "Configure gating via alumniManagement/rehireEligibility features."),
        "Rehire": fmt(F, "UI: Rehire button", "POST /v1/admin/lifecycle/alumni/:id/rehire", "Panel reload after action.", "Feature gated via rehireEligibility."),
        "Blacklist": fmt(F, "UI: Blacklist button (reason prompt)", "POST /v1/admin/lifecycle/alumni/:id/blacklist", "Panel reload after action.", "Feature gated via rehireEligibility."),
        "Timeline": fmt(P, "UI: Recent Activity list acts as timeline (no stage timeline viz)", "workflow.activityFeed", "Realtime: socket refresh.", "N/A"),
        "Reports": fmt(P, "UI: minimal analytics via dashboard.analytics + full reports in LifecycleReportsHub", "GET /dashboard + /analytics; GET /v1/admin/lifecycle/reports/*", "No socket inside reports hub.", "Reports accessible via EmployeeLifecycle query reports=1."),
        "Forms": fmt(P, "UI: exit interview questionnaires are the primary form UI in offboarding hub", "Exit interview form GET /forms/:formDefinitionId", "Panel reload on draft/submit.", "Configure gating via exitInterview feature."),
        "Workflow": fmt(F, "UI: create/selection/actions for offboarding workflows", "POST /v1/admin/onboarding/offboarding; POST /actions/:action; DELETE /offboarding/:id", "Realtime: socket for changes.", "Configure gating via lifecycle feature registry."),
        "Activity Logs": fmt(F, "UI: activityFeed list + HR Internal Notes section", "activityFeed in payload; POST /v1/admin/onboarding/offboarding/:wfId/notes", "Realtime: socket refresh and panel reload.", "N/A"),
        "Notifications": fmt(P, "UI: notification feed not shown in hub; enterprise notification center exists", "GET/PATCH /v1/admin/lifecycle/notifications/center; other notification APIs in /v1/admin/notifications/*", "N/A", "Configure gating via notificationCenter feature."),
        "Configure": fmt(F, "UI: LifecycleConfigurationPanel controls offboarding feature registry toggles", "GET /v1/admin/lifecycle/registry; PATCH /v1/admin/lifecycle/settings", "Soft refresh after save.", "Configure is core integration."),
        "Enterprise": fmt(F, "UI: LifecycleEnterpriseHub modules and bulk/notifications search", "GET/POST /v1/admin/lifecycle/* per tab", "Manual refresh.", "Feature gated tabs."),
        "Progress Tracking": fmt(F, "UI: progress rings and breakdown for tasks/checklist/assets/knowledge/access", "breakdown derived from workflow payload", "Socket refresh.", "N/A"),
        "Status Cards": fmt(F, "UI: KPI stat cards + workflow status badges + progress rings", "dashboard.stats, workflow.status", "Socket refresh.", "N/A"),
        "Buttons": fmt(F, "UI: workflow controls, checklist toggles, action buttons, letters download/upload", "Various PUT/POST endpoints.", "Refresh after actions.", "Feature gated visibility."),
        "Tables": fmt(M, "OffboardingHub is card/list oriented (limited tables).", "N/A", "N/A", "N/A"),
        "Empty States": fmt(F, "UI: Empty messages per panel (no workflow, no letters, no alumni).", "N/A", "N/A", "N/A"),
        "Loading States": fmt(F, "UI: LoadingSpinner per hub/panels.", "N/A", "N/A", "N/A"),
    }

    off_rows_table = []
    for feat in off_features:
        off_rows_table.append([feat, off[feat]])
    add_table(doc, headers=["Feature", "Admin/HR Offboarding (OffboardingHub)"], rows=off_rows_table)

    add_heading(doc, "Step 5 - Identified UI components likely missing vs your design", level=1)
    add_para(doc, "1. Rich stage timeline visualization (beyond status chips + progress rings).")
    add_para(doc, "2. Unified employee profile sidebar embedded in onboarding/offboarding detail (with consistent status badges and quick actions).")
    add_para(doc, "3. In-hub analytics charts (charts in OnboardingHub/OffboardingHub rather than only in LifecycleReportsHub).")
    add_para(doc, "4. Hub-level notifications feed + action center UI.")
    add_para(doc, "5. Approval queue panels for onboarding/offboarding items (forms/tasks/signatures/clearance approval workflows).")
    add_para(doc, "6. Document preview cards embedded in detail (quick preview/download/verify).")
    add_para(doc, "7. Better approval/error UX states for F&F and access revocation actions.")

    add_heading(doc, "Step 6 - Verification risks (static evidence)", level=1)
    add_para(doc, "- Reporting UI risk: `LifecycleReportsHub` generates synthetic/fallback values for chart rendering when trend data is missing (may violate strict “no dummy statistics”).")
    add_para(doc, "- Notifications UI risk: enterprise config exists, but hubs do not render a notifications feed.")
    add_para(doc, "- Public pre-boarding portal risk: `server/routes/Public/LifecyclePortalRoutes.js` supports only `/offer/:token` accept/decline; no document/form upload endpoints.")

    doc.save(str(OUT_PATH))
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()

