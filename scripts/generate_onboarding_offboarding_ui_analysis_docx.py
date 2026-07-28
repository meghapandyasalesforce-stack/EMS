from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt


OUT_PATH = Path(__file__).resolve().parents[1] / "Onboarding_Offboarding_Modules_Existing_Implementation_Analysis.docx"


F = "Fully Functional"
P = "Partially Implemented"
M = "Not Implemented / Missing"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for p in h.runs:
        p.font.size = Pt(14 if level == 1 else 12)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_table(doc: Document, headers, rows) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = str(val) if val is not None else ""


def main() -> None:
    doc = Document()

    # Title
    title = doc.add_paragraph("EMS - Admin & HR Onboarding/Offboarding Analysis (Pre-UI Redesign)")
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(16)
    add_para(doc, f"Generated: {date.today().isoformat()}")
    add_para(
        doc,
        "Method: static code-path tracing across the existing React UI components and the backend endpoints they call. "
        "Realtime behavior was inferred from socket.io usage; configure integration was inferred from lifecycle feature-gating (useLifecycleFeatures + isFeatureEnabled).",
    )

    # Step 1
    add_heading(doc, "Step 1 - What exists today (Current modules & entrypoints)", level=1)
    add_para(
        doc,
        "Onboarding (Admin/HR lifecycle hub): `vite_client/src/HR/Onboarding/OnboardingHub.jsx` "
        "with sub-panels: OnboardingBgvPanel, OnboardingFormsPanel, OnboardingGoalsPanel, OnboardingMilestonesPanel, "
        "DocumentHub, ProbationCenter, ItProvisioningPanel, LmsDashboard.",
    )
    add_para(
        doc,
        "Offboarding (Admin/HR lifecycle hub): `vite_client/src/HR/Offboarding/OffboardingHub.jsx` "
        "with sub-panels: ExitInterviewWorkflowPanel, OffboardingLettersPanel, OffboardingAlumniPanel.",
    )
    add_para(
        doc,
        "Enterprise/Configure/Reports: `vite_client/src/HR/Onboarding/EmployeeLifecycle.jsx` "
        "which routes Admin/HR users to `LifecycleEnterpriseHub`, `LifecycleConfigurationPanel`, and `LifecycleReportsHub`.",
    )
    add_para(
        doc,
        "Admin Onboarding legacy management UI (separate from lifecycle hub): "
        "`vite_client/src/HR/Onboarding/OnboardingPage.jsx` and `OnboardingDetailPage.jsx` (admin onboarding CRUD).",
    )

    # Step 2
    add_heading(doc, "Step 2 - Feature coverage (Fully / Partially / Missing)", level=1)
    add_heading(doc, "Onboarding Feature Coverage Matrix", level=2)

    onb_features = [
        "Dashboard",
        "Statistics Cards",
        "Workflow Stages",
        "Employee Details",
        "Search",
        "Filters",
        "Documents",
        "Document Hub",
        "Forms",
        "Form Builder",
        "Tasks",
        "Checklist",
        "Background Verification",
        "Buddy",
        "Goals",
        "Probation",
        "30-60-90 Plan",
        "Software Requests",
        "IT Assets",
        "Training/LMS",
        "Timeline",
        "Activity Log",
        "Notes",
        "Comments",
        "Approvals",
        "Notifications",
        "Reports",
        "Workflow Builder",
        "Offer Letter",
        "Candidate Details",
        "Configure",
        "Enterprise",
        "Progress Tracking",
        "Status Cards",
        "Buttons",
        "Tables",
        "Empty States",
        "Loading States",
    ]

    admin_onb = {
        "Dashboard": (P, "UI: legacy workflow cards/list; no KPI dashboard."),
        "Statistics Cards": (M, "Not shown in legacy admin onboarding pages."),
        "Workflow Stages": (P, "UI: status chips/filter + status badges."),
        "Employee Details": (P, "UI: detail page has start/expected completion + assigned HR/buddy + welcome."),
        "Search": (F, "UI: search by employee name/email."),
        "Filters": (F, "UI: statusFilter dropdown."),
        "Documents": (M, "No documents section in legacy admin onboarding UI."),
        "Document Hub": (M, "Not surfaced in legacy admin onboarding UI."),
        "Forms": (M, "No mandatory forms UI in legacy admin onboarding."),
        "Form Builder": (M, "Configure-only area via FormBuilderHub."),
        "Tasks": (F, "UI: tasks list + start/complete + task status updates via API."),
        "Checklist": (M, "No checklist UI in legacy admin onboarding."),
        "Background Verification": (P, "UI: BGV case + HR check status dropdowns; not the full HR BGV workflow UI."),
        "Buddy": (P, "UI: buddy displayed in sidebar; no buddy check-in/meeting workflow UI."),
        "Goals": (M, "No goals/probation/KRA UI in legacy admin onboarding."),
        "Probation": (M, "No probation UI in legacy admin onboarding."),
        "30-60-90 Plan": (M, "No milestones/day plan UI in legacy admin onboarding."),
        "Software Requests": (M, "No IT/software request UI in legacy admin onboarding."),
        "IT Assets": (M, "No asset UI in legacy admin onboarding."),
        "Training/LMS": (M, "No training/LMS UI in legacy admin onboarding."),
        "Timeline": (M, "No structured timeline visualization in legacy admin onboarding."),
        "Activity Log": (M, "No activity feed/timeline in legacy admin onboarding."),
        "Notes": (M, "No notes UI in legacy admin onboarding."),
        "Comments": (M, "No comment-thread UI found."),
        "Approvals": (M, "No onboarding approval queue UI in legacy pages."),
        "Notifications": (M, "No notifications UI in legacy admin onboarding pages."),
        "Reports": (M, "Reports are in LifecycleReportsHub, not legacy onboarding pages."),
        "Workflow Builder": (M, "Workflow builder only in Configure area via WorkflowBuilder."),
        "Offer Letter": (M, "Offer management exists in a separate offer hub (OfferTemplateManager/OfferLetterPanel)."),
        "Candidate Details": (P, "Legacy onboarding UI shows employee details; no offer-candidate acceptance flow."),
        "Configure": (M, "Not present in legacy onboarding UI."),
        "Enterprise": (M, "Not present in legacy onboarding UI."),
        "Progress Tracking": (P, "UI: progress bars and completed/remaining derived from task list."),
        "Status Cards": (P, "UI: status badges on cards."),
        "Buttons": (F, "UI buttons/modals exist for CRUD and task updates."),
        "Tables": (M, "Legacy UI is card/grid based, not table/grid records."),
        "Empty States": (F, "Empty states present for list/tasks."),
        "Loading States": (F, "LoadingSpinner used."),
    }

    hr_onb = {
        "Dashboard": (F, "UI: OnboardingHub candidate/workflow table + KPI strip when no workflow selected."),
        "Statistics Cards": (F, "UI: CmStat KPIs (pending actions, offer acceptance rate, avg onboarding time)."),
        "Workflow Stages": (F, "UI: statusFilter chips + status badges."),
        "Employee Details": (F, "UI: selected workflow header includes HR/manager/buddy + welcome message."),
        "Search": (F, "UI: CmSearch filters returned workflows client-side."),
        "Filters": (F, "UI: status chips."),
        "Documents": (P, "UI: documents section shows status + navigation; uploads/verification live in Documents area."),
        "Document Hub": (F, "UI: DocumentHub provides centralized lifecycle document verification/expiry/compliance."),
        "Forms": (F, "UI: OnboardingFormsPanel (mandatory forms + submit flow)."),
        "Form Builder": (F, "UI: FormBuilderHub provides definition builder + field toolbox."),
        "Tasks": (F, "UI: tasks by role; start/continue/mark complete with blockers enforcement."),
        "Checklist": (F, "UI: checklist checkboxes with toggle + backend updates."),
        "Background Verification": (F, "UI: OnboardingBgvPanel with HR dashboard + checks UI; config can disable BGV."),
        "Buddy": (P, "UI: buddy check-ins update status; automation/matching and welcome-kit UX not surfaced."),
        "Goals": (F, "UI: OnboardingGoalsPanel loads/assigns probation goals & KRA items."),
        "Probation": (F, "UI: ProbationCenter for pipeline, extend, and review sign-off."),
        "30-60-90 Plan": (F, "UI: OnboardingMilestonesPanel for day-plan reviews."),
        "Software Requests": (F, "UI: ItProvisioningPanel for workflow-scoped software requests + approve/provision."),
        "IT Assets": (F, "UI: Equipment & Assets section updates asset status for HR view."),
        "Training/LMS": (F, "UI: training modules list + LmsDashboard (assign/sync + certifications)."),
        "Timeline": (P, "UI: recent activity list exists; not a full graphical timeline component."),
        "Activity Log": (F, "UI: activityFeed list rendered as Recent Activity."),
        "Notes": (F, "UI: HR-only Internal Notes section."),
        "Comments": (M, "No comment threads UI; only internal notes."),
        "Approvals": (P, "Offers and some policy signing use approval/submit workflows; onboarding hub itself lacks approval queues."),
        "Notifications": (P, "Enterprise notification center is configurable; no dedicated notification feed inside the hub."),
        "Reports": (P, "Analytics KPI strip in hub; deep drill-down in LifecycleReportsHub."),
        "Workflow Builder": (F, "Configure via WorkflowBuilder supports onboarding/offboarding stage design."),
        "Offer Letter": (F, "Offer templates + offer generation and candidate portal accept/decline."),
        "Candidate Details": (P, "Hub shows candidate/workflow progression; offer-candidate acceptance workflow is separate."),
        "Configure": (F, "UI: LifecycleConfigurationPanel feature toggles + policies + approvals/escalations/reminders."),
        "Enterprise": (F, "UI: LifecycleEnterpriseHub provides command center modules + enterprise tabs."),
        "Progress Tracking": (F, "UI: progress bars and progress rings across tasks/documents/checklist/training/policies/assets."),
        "Status Cards": (F, "UI: KPI cards + workflow status badges."),
        "Buttons": (F, "UI: action buttons for workflow/task/checklist/policy signing/etc."),
        "Tables": (F, "UI: cm-table workflow candidate list."),
        "Empty States": (F, "UI: CmEmpty/empty prompts exist per module."),
        "Loading States": (F, "LoadingSpinner/CM loading states exist across panels."),
    }

    rows = []
    for feat in onb_features:
        a_sym, a_note = admin_onb.get(feat, (M, ""))
        b_sym, b_note = hr_onb.get(feat, (M, ""))
        rows.append([feat, f"{a_sym}\\n{a_note}", f"{b_sym}\\n{b_note}"])

    add_table(doc, headers=["Feature", "Admin Onboarding (Legacy UI)", "HR Onboarding (Lifecycle Hub)"], rows=rows)

    add_heading(doc, "Offboarding Feature Coverage Matrix (Admin + HR use OffboardingHub)", level=2)

    off_features = [
        "Dashboard",
        "Statistics Cards",
        "Exit Workflow",
        "Notice Period",
        "Exit Interview",
        "Clearance",
        "Knowledge Transfer",
        "Asset Return",
        "Access Revocation",
        "F&F Settlement",
        "Payroll Bridge",
        "Generated Letters",
        "Experience Letter",
        "Relieving Letter",
        "Service Certificate",
        "Alumni",
        "Rehire",
        "Blacklist",
        "Timeline",
        "Reports",
        "Forms",
        "Workflow",
        "Activity Logs",
        "Notifications",
        "Configure",
        "Enterprise",
        "Progress Tracking",
        "Status Cards",
        "Buttons",
        "Tables",
        "Empty States",
        "Loading States",
    ]

    off = {
        "Dashboard": (F, "UI: OffboardingHub header + stats grid when no workflow selected."),
        "Statistics Cards": (F, "UI: KPI stats like Total Separations, Active, Upcoming, Overdue, pending assets/clearances."),
        "Exit Workflow": (F, "UI: workflow list + selected workflow detail + publish/pause/resume/cancel actions."),
        "Notice Period": (F, "UI: noticePeriod card + HR buyout/leave adjustment actions (feature gated)."),
        "Exit Interview": (F, "UI: ExitInterviewWorkflowPanel sessions UI; provision, draft save, submit, close."),
        "Clearance": (F, "UI: Clearance checklist checkboxes (gated by clearance/hrClearance)."),
        "Knowledge Transfer": (F, "UI: knowledge transfer items with status dropdown."),
        "Asset Return": (F, "UI: Asset Return section with asset status select (gated by assetReturn)."),
        "Access Revocation": (F, "UI: Access Revocation section + auto-revoke action (gated)."),
        "F&F Settlement": (P, "UI: calculate components + submit for approval + PDF + notes; may still rely on manual inputs/statutory gaps."),
        "Payroll Bridge": (P, "UI: create payroll run action exists; hub lacks detailed payroll-run status view."),
        "Generated Letters": (F, "UI: generated letters list with download links."),
        "Experience Letter": (F, "UI: dedicated Experience Letter upload/replace + view/download."),
        "Relieving Letter": (P, "UI: surfaced only via generic generated letters list; no dedicated rel./cert cards."),
        "Service Certificate": (P, "UI: surfaced only via generic generated letters list."),
        "Alumni": (F, "UI: alumni database panel (HR-only) with KPI mini-cards."),
        "Rehire": (F, "UI: rehire button triggers alumni rehire endpoint."),
        "Blacklist": (F, "UI: blacklist button prompts for reason and triggers blacklist endpoint."),
        "Timeline": (P, "UI: Recent Activity list exists; no full timeline component."),
        "Reports": (P, "UI: basic analytics in hub; full drill-down in LifecycleReportsHub."),
        "Forms": (P, "UI: exit interview questionnaires are forms; no broader offboarding form engine UI in hub."),
        "Workflow": (F, "UI: lifecycle workflow progression + actions."),
        "Activity Logs": (F, "UI: activityFeed list + HR internal notes."),
        "Notifications": (P, "Enterprise notification center is configurable; hub lacks a notifications feed."),
        "Configure": (F, "UI: LifecycleConfigurationPanel toggles offboarding features/policies."),
        "Enterprise": (F, "UI: LifecycleEnterpriseHub command center modules."),
        "Progress Tracking": (F, "UI: progress rings + status badges + breakdown values for tasks/checklist/assets/knowledge/access."),
        "Status Cards": (F, "UI: KPI cards + status badges."),
        "Buttons": (F, "UI: action buttons/selects exist."),
        "Tables": (M, "UI is card/list based; limited table usage."),
        "Empty States": (F, "UI empty prompts exist per section."),
        "Loading States": (F, "LoadingSpinner used across hub/panels."),
    }

    rows2 = []
    for feat in off_features:
        sym, note = off.get(feat, (M, ""))
        rows2.append([feat, f"{sym}\\n{note}"])

    add_table(doc, headers=["Feature", "Admin/HR Offboarding (OffboardingHub)"], rows=rows2)

    # Step 3/4/5
    add_heading(doc, "Step 3 - Comparison with uploaded design (best-effort mapping)", level=1)
    add_para(
        doc,
        "The repository does not contain readable UI mock assets for the “uploaded design”. "
        "This comparison uses the section targets described in your prompt (rich dashboard, better timeline, modern cards, etc.) "
        "and maps them to the current implementation quality.",
    )
    comp_headers = ["Module", "Section", "Current", "Uploaded Design", "Required Changes"]
    comp_rows = [
        ["Onboarding", "Dashboard", "KPI strip + workflow table; limited charting", "Rich Dashboard UI", "UI redesign: unify KPI + add chart panels + align card/grid layout."],
        ["Onboarding", "Workflow", "Status chips + progress rings + tasks list", "Better Timeline / Workflow Visualization", "UI redesign: introduce stage timeline visualization and per-stage action panels."],
        ["Onboarding", "Documents", "Documents status + navigation; separate DocumentHub experience", "Modern Cards + Preview", "UI redesign: embed document preview/verification cards inside onboarding/offboarding detail."],
        ["Onboarding", "Forms", "Modal-based mandatory forms list + submit flow", "Better Layout + Review flow", "UI redesign: simplify form UX with inline progress, review/approval states, and consistent buttons."],
        ["Onboarding", "BGV", "Dedicated BGV panel with checks UI", "Regulated-industry BGV Dashboard", "UI redesign: standardize BGV checklist visuals, add right-side timeline, and consistent status badges."],
        ["Offboarding", "Exit Workflow", "Workflow cards + detail actions", "Modern Timeline UI", "UI redesign: add richer stage timeline + role-aware action panels."],
        ["Offboarding", "Clearance", "Checkbox checklist; feature-gated", "Better Cards + Summary", "UI redesign: add clearance summary cards and grouped sections for assets/clearances."],
        ["Offboarding", "F&F", "Calculator + submit approval + PDF + payroll run button", "Better Dashboard/Analytics", "UI redesign: add component breakdown charts and better approval/error states."],
        ["Both", "Activity / Timeline", "Recent Activity list + internal notes", "Rich Timeline + Activity Feed", "UI redesign: convert activity list to true timeline with filters and grouping."],
    ]
    add_table(doc, comp_headers, comp_rows)

    add_heading(doc, "Step 4 - Reuse existing features (no backend/API rewrite)", level=1)
    add_para(doc, "Recommended reuse candidates: `OnboardingHub`, `OffboardingHub`, task/checklist update handlers, `OnboardingBgvPanel`, `OnboardingFormsPanel`, `OnboardingGoalsPanel`, `OnboardingMilestonesPanel`, `DocumentHub`, `ProbationCenter`, `ItProvisioningPanel`, `LmsDashboard`, `WorkflowBuilder`, `FormBuilderHub`, `LetterGeneratorHub`, `OfferTemplateManager`/`OfferLetterPanel`, `LifecycleConfigurationPanel`, `LifecycleEnterpriseHub`, `LifecycleReportsHub`.")
    add_para(doc, "For Admin Onboarding legacy pages, the safest approach is to redesign UI to reuse the lifecycle hub/panels rather than maintaining a parallel incomplete UI surface.")

    add_heading(doc, "Step 5 - Missing UI components likely required by the design", level=1)
    missing_components = [
        "Rich workflow stage timeline visualization (not just status chips/progress rings).",
        "Unified employee profile sidebar (HR/Admin controls + key status badges) inside onboarding/offboarding detail.",
        "Hub-level analytics charts (not only KPI strips or only in LifecycleReportsHub).",
        "In-hub notification feed (instead of only enterprise notification configuration).",
        "Approval panels/queues UI for onboarding/offboarding forms and policy workflows.",
        "Document preview cards and quick actions embedded into onboarding/offboarding detail views.",
        "Consistent approval/status action panels across tasks, checklists, assets, knowledge transfer, and letters.",
    ]
    for m in missing_components:
        add_para(doc, f"- {m}")

    add_heading(doc, "Step 6 - Functionality verification (static checks + risks)", level=1)
    add_para(doc, "- Realtime: onboarding/offboarding hubs trigger soft refresh via socket.io events (`onboarding:changed`, `offboarding:changed`). Sub-panels mostly rely on reload after actions (not always socket-driven).")
    add_para(doc, "- Reporting UI risk: LifecycleReportsHub includes fallback/synthetic chart values for rendering when trend data is missing (e.g., “fallback sample-style rows” and sparkline generation).")
    add_para(doc, "- Notifications UI risk: lifecycle notification configuration exists (`/admin/lifecycle/notifications/center`), but the hubs do not render a dedicated notifications feed.")
    add_para(doc, "- Pre-boarding portal gap: public `/public/lifecycle/offer/:token` route supports only accept/decline; there are no public endpoints for candidate document/form upload beyond offer response.")

    add_heading(doc, "Step 7 - Final Analysis Summary (before UI redesign)", level=1)
    add_para(doc, "Onboarding: HR lifecycle hub is broad and feature-rich (tasks/checklists/BGV/forms/goals/probation/30-60-90/software+IT/training/document hub/activity feed/notes). Admin onboarding legacy UI is largely incomplete compared to the lifecycle hub (mostly workflows + tasks + limited BGV).")
    add_para(doc, "Offboarding: OffboardingHub provides a comprehensive role-based UI surface (exit workflow, checklist/clearance, knowledge transfer, assets, access revocation, notice period, F&F actions, exit interviews, letters, alumni actions, timeline + notes). Some deeper UX expectations (timeline visualization, approvals UX, hub-level notifications feed, detailed payroll-run status) are partial.")
    add_para(doc, "Next step: once you provide the actual UI mock assets for the uploaded design, we can finalize the UI component-by-component redesign plan while preserving all business logic, APIs, RBAC, feature gating, and realtime behavior.")

    doc.save(str(OUT_PATH))
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()

