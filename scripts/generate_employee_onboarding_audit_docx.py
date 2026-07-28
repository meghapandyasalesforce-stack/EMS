"""
Employee Onboarding discovery audit — documentation only (what exists today).
Output: Employee_Onboarding_Audit.docx at repo root.
"""
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_PATH = Path(__file__).resolve().parents[1] / "Employee_Onboarding_Audit.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for p in h.runs:
        p.font.size = Pt(14 if level == 1 else 12)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_table(doc: Document, headers, rows) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = str(val) if val is not None else ""


def main() -> None:
    doc = Document()
    title = doc.add_paragraph("PeopleFlow / EMS — Employee Onboarding Audit (As-Built)")
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(16)
    add_para(
        doc,
        f"Generated: {date.today().isoformat()}. Scope: employee lifecycle onboarding (multi-tenant via "
        "organizationId). Discovery from server models, routes, services, and vite_client — no new features proposed.",
    )

    # ── 1. Data Model ──
    add_heading(doc, "1. Data Model", level=1)
    add_para(
        doc,
        "Primary persistence lives in server/models/onboarding.model.js (OnboardingWorkflow, OnboardingTemplate, "
        "plus offboarding models in the same file). Lifecycle configuration, forms, and BGV use "
        "server/models/lifecycle.model.js and server/models/lifecycleGap.model.js. Canonical employee uploads use "
        "server/models/document.model.js (Document). Tenant scoping: organizationId is required and indexed on "
        "OnboardingWorkflow, OnboardingTemplate, LifecycleSettings, LifecycleFormDefinition, LifecycleFormResponse, "
        "BackgroundVerificationCase, and Document.",
    )

    add_heading(doc, "1.1 OnboardingWorkflow (employee instance)", level=2)
    add_table(
        doc,
        ["Field", "Type", "Purpose"],
        [
            ("organizationId", "ObjectId → Organization", "Tenant discriminator (required, indexed)"),
            ("employeeId", "ObjectId → Employee", "Subject of onboarding (required, indexed)"),
            ("templateId", "ObjectId → OnboardingTemplate", "Source template"),
            ("backgroundVerificationCaseId", "ObjectId → BackgroundVerificationCase", "Linked BGV case"),
            ("startDate / joiningDate / scheduledStartDate", "Date", "Workflow timing; scheduled start from org settings"),
            ("expectedCompletionDate / actualCompletionDate", "Date", "SLA target and completion"),
            ("status", "String enum", "Draft | Scheduled | Not Started | In Progress | Completed | Paused | Cancelled | Overdue | On Hold"),
            ("pausedFromStatus", "String enum", "Prior status when Paused"),
            ("tasks[]", "Subdoc", "Role-assigned tasks (see 1.2)"),
            ("documents[]", "Subdoc", "Required upload slots synced to Document model (see 1.3)"),
            ("checklist[]", "Subdoc", "Joining checklist items by role"),
            ("progress", "Number 0–100", "Overall completion percentage"),
            ("progressBreakdown", "Object", "Per-dimension %: tasks, documents, checklist, training, policies, assets, forms, goals, bgv, software, byRole"),
            ("bgvProgress / bgvStatus / bgvBlocksJoining", "Number + enum + Boolean", "Mirrored BGV state; gates completion when blockJoiningUntilClear"),
            ("assets[]", "Subdoc", "IT/HR physical assets"),
            ("trainingModules[]", "Subdoc", "Video/quiz/document training"),
            ("policySignatures[]", "Subdoc", "E-sign policy pack"),
            ("buddyCheckins[]", "Subdoc", "Buddy program milestones"),
            ("managerActions[]", "Subdoc", "Manager onboarding actions"),
            ("milestones[]", "Subdoc", "30/60/90 day reviews"),
            ("probation", "Object", "Probation window and reviews"),
            ("softwareRequests[]", "Subdoc", "IT access provisioning tickets"),
            ("formsProgress / mandatoryFormsCompleted / formGateStatus", "Number, Boolean, Mixed", "Digital forms gate for workflow completion"),
            ("goalsProgress / assignedGoalIds / onboardingGoals", "Number, refs, Mixed", "Probation goals integration"),
            ("policyAccepted / policyAcceptedAt", "Boolean, Date", "Legacy bulk policy accept flag"),
            ("activityFeed / notes / notificationHistory", "Arrays", "Audit trail and HR notes (visibility on notes)"),
            ("assignedHR / assignedBuddy / assignedManager", "ObjectId refs", "Workflow owners"),
            ("isTemplate", "Boolean", "false for employee workflows"),
        ],
    )
    add_para(doc, "Note: OnboardingTemplate defines stages[] (lifecycleStageSchema), but OnboardingWorkflow schema does not declare a stages[] field. Runtime stage code (lifecycleWorkflowRuntimeService) sets workflow.stages in memory; initializeOnboardingStages is never called from workflow creation.")

    add_heading(doc, "1.2 onboardingTask (embedded)", level=2)
    add_table(
        doc,
        ["Field", "Type", "Purpose / enums"],
        [
            ("title, description", "String", "Task label and help text"),
            ("category", "enum", "Documentation | IT Setup | HR Paperwork | Training | Orientation | Equipment | Access | Other"),
            ("assignedTo", "ObjectId → Employee", "Concrete assignee"),
            ("assigneeRole", "enum", "employee | hr | manager | buddy | it | payroll | admin | other"),
            ("dueDate, startedDate, completedDate, overdueAt", "Date", "Scheduling and SLA"),
            ("status", "enum", "Pending | In Progress | Waiting | Blocked | Completed | Cancelled | Overdue"),
            ("priority", "enum", "Low | Medium | High | Critical"),
            ("mandatory", "Boolean", "Counts toward progress when feature enabled"),
            ("phase", "String", "Timeline bucket (e.g. Day 1, Week 1)"),
            ("dependencies[]", "ObjectId", "Other task IDs that must be Completed first"),
            ("requiredDocuments[]", "String", "Document type keys required before completion"),
            ("requiresAllMandatoryDocsVerified", "Boolean", "HR verification task gate"),
            ("attachments, comments, completionNotes, notes", "Various", "Collaboration and evidence"),
        ],
    )

    add_heading(doc, "1.3 onboardingDocument (embedded)", level=2)
    add_table(
        doc,
        ["Field", "Type", "Purpose / enums"],
        [
            ("canonicalDocumentId", "ObjectId → Document", "Link to canonical employee document library"),
            ("documentType", "enum (large)", "AADHAAR, PAN, PASSPORT, DEGREE, EXPERIENCE, RESUME, OFFER_ACK, PHOTO, BANK_DETAILS, OTHER, plus legacy labels"),
            ("status", "enum", "Pending | Uploaded | Verified | Rejected"),
            ("mandatory", "Boolean", "Required for progress / HR verify task"),
            ("uploadedBy, verifiedBy, verifiedAt, rejectionReason", "Refs / dates / string", "Verification workflow"),
            ("version, versionHistory[]", "Number, array", "Re-upload history"),
        ],
    )

    add_heading(doc, "1.4 Other embedded workflow collections", level=2)
    add_table(
        doc,
        ["Collection", "Key status enums", "Notes"],
        [
            ("checklist[]", "completed: Boolean", "assignedRole enum; featureKey for gating"),
            ("assets[]", "Pending | Assigned | Returned | Lost | Damaged", "assetType enum (Laptop, ID Card, …)"),
            ("trainingModules[]", "Not Started | In Progress | Completed | Failed", "type: video | quiz | document"),
            ("policySignatures[]", "Pending | Signed", "documentType: OFFER_LETTER, NDA, COMPANY_POLICIES, …"),
            ("buddyCheckins[]", "Pending | Scheduled | Completed", "type: welcome_meeting | first_week | day_30"),
            ("managerActions[]", "Pending | In Progress | Completed", "probation_goals, first_project, one_on_one, review, …"),
            ("milestones[]", "Pending | In Progress | Completed | Overdue | Waived", "day: 30 | 60 | 90; employee/manager/hr reviews"),
            ("probation", "Not Started | Active | Extended | Confirmed | Rejected | Terminated | Waived", "midReview / finalReview sub-objects"),
            ("softwareRequests[]", "status + approvalStatus enums", "Provisioning lifecycle; featureKey optional"),
        ],
    )

    add_heading(doc, "1.5 OnboardingTemplate", level=2)
    add_para(
        doc,
        "organizationId (required). Defines name, expectedDays, version/status (Draft | Published | Archived), "
        "stageGroups[], stages[] (lifecycleStageSchema: key, name, order, mandatory, featureKey, dependsOn[], dueDaysFromStart, …), "
        "and template arrays for tasks, checklist, documents, assets, trainingModules, policySignatures.",
    )

    add_heading(doc, "1.6 LifecycleSettings (per org)", level=2)
    add_para(
        doc,
        "organizationId unique. modules.onboarding.enabled, enabledFeatures map (feature flags), defaults.onboardingTemplateId, "
        "onboarding.autoCreateOnEmployeeRegister, bgvPolicy.blockJoiningUntilClear, documentOverrides, reminders, escalations, branding.",
    )

    add_heading(doc, "1.7 LifecycleFormDefinition / LifecycleFormResponse", level=2)
    add_table(
        doc,
        ["Model", "Key fields", "Status enums"],
        [
            ("LifecycleFormDefinition", "organizationId, name, category (onboarding, employee_profile, …), fields[], featureKey, stageKey", "Draft | Published | Archived"),
            ("LifecycleFormResponse", "organizationId, formDefinitionId, employeeId, workflowId, workflowType, responses (Mixed)", "Draft | Submitted | Pending Approval | Approved | Rejected"),
        ],
    )

    add_heading(doc, "1.8 BackgroundVerificationCase (lifecycleGap.model.js)", level=2)
    add_para(
        doc,
        "organizationId, employeeId, onboardingWorkflowId, checks[] (checkType: employment | education | address | reference | criminal | identity | other; "
        "status includes Not Started, Pending Employee Submission, Draft Saved, Submitted, Under HR Review, Clear, Failed, …). "
        "Linked from workflow.backgroundVerificationCaseId when backgroundVerification feature is enabled at creation.",
    )

    add_heading(doc, "1.9 Document (canonical library)", level=2)
    add_para(
        doc,
        "organizationId, uploadedBy, fileUrl, approvalStatus, employeeJobDocumentType, tags — synced into workflow.documents via onboardingService.synchronizeWorkflowDocuments.",
    )

    add_heading(doc, "1.10 Organization onboarding (distinct product surface)", level=2)
    add_para(
        doc,
        "Tenant SaaS setup wizard (organization.model onboardingSteps) — NOT employee onboarding. Steps defined in onboardingDefaults.ORG_ONBOARDING_STEPS (welcome, email, company-info, … complete).",
    )

    # ── 2. Backend APIs ──
    add_heading(doc, "2. Backend APIs (Employee-facing)", level=1)
    add_para(
        doc,
        "Base paths: /api/v1/both/onboarding (isAuthenticatedBoth: JWT role admin or employee), "
        "/api/v1/both/lifecycle/employee/* (isAuthenticated: requires active Employee profile; admins with linked profile may call). "
        "Feature gates: requireLifecycleModule('onboarding') and requireLifecycleFeature(featureKey) on selected lifecycle routes. "
        "HR-only admin routes mount at /api/v1/admin/onboarding and /api/v1/admin/lifecycle.",
    )

    add_heading(doc, "2.1 /api/v1/both/onboarding", level=2)
    add_table(
        doc,
        ["Method", "Path", "Purpose", "Request body (summary)", "Response", "Auth / guards"],
        [
            ("GET", "/my-workflow", "Load caller's onboarding workflow", "—", "{ success, data: populated OnboardingWorkflow (sanitized) }", "isAuthenticatedBoth"),
            ("GET", "/my-tasks", "Open tasks for caller across workflows", "—", "{ success, data: [{ workflow, task }] }", "isAuthenticatedBoth"),
            ("GET", "/dashboard", "Role-aware dashboard (employee self, manager team, or HR stats if admin)", "Query: view", "{ success, data: { workflows, stats, role, templates? } }", "isAuthenticatedBoth"),
            ("GET", "/workflow/:workflowId/bgv", "BGV case for workflow", "—", "BGV payload or empty if feature disabled (allowEmpty)", "isAuthenticatedBoth + requireLifecycleFeature(backgroundVerification, allowEmpty GET)"),
            ("POST", "/workflow/:workflowId/bgv", "Create/link BGV from workflow", "Case options", "BGV case", "isAuthenticatedBoth + backgroundVerification"),
            ("PATCH", "/workflow/:workflowId/bgv/checks/:checkId", "Update check from onboarding context", "check fields", "Updated check", "isAuthenticatedBoth + backgroundVerification"),
            ("PUT", "/workflow/:id/tasks/:taskId", "Update task status/notes", "{ status: Pending|In Progress|Completed, … }", "Populated workflow", "isAuthenticatedBoth; employee assignee or workflow employee for employee-role tasks"),
            ("PUT", "/workflow/:id/checklist/:itemId", "{ completed }", "Checklist toggle", "Populated workflow", "isAuthenticatedBoth; role-based"),
            ("POST", "/workflow/:id/accept-policy", "Set policyAccepted flag", "—", "Workflow", "isAuthenticatedBoth"),
            ("POST", "/workflow/:id/signatures/:signatureId/sign", "Sign policy pack item", "{ signatureData | acknowledgment }", "Workflow", "isAuthenticatedBoth; workflow employee only; eSignature feature"),
            ("PUT", "/workflow/:id/training/:trainingId", "Update training module", "{ status, score, answers }", "Workflow", "isAuthenticatedBoth; employee or HR; mandatoryTraining/trainingLms feature"),
            ("PUT", "/workflow/:id/buddy-checkins/:checkinId", "Buddy check-in", "status, feedback, rating", "Workflow", "isAuthenticatedBoth; assigned buddy or HR; buddyProgram feature"),
            ("PUT", "/workflow/:id/manager-actions/:actionId", "Manager action item", "status, notes", "Workflow", "isAuthenticatedBoth; manager/HR"),
        ],
    )
    add_para(doc, "No PUT /workflow/:id/assets/:assetId on both/onboarding — asset updates use lifecycle employee route below.")

    add_heading(doc, "2.2 /api/v1/both/lifecycle/employee (onboarding-related)", level=2)
    add_table(
        doc,
        ["Method", "Path", "Purpose", "Body", "Response", "Auth / guards"],
        [
            ("GET", "/employee/onboarding", "Alternate full workflow fetch", "—", "{ success, data: workflow + employeeTasks filter }", "isAuthenticated + requireLifecycleModule(onboarding)"),
            ("GET", "/employee/dashboard", "Lifecycle summary dashboard", "—", "onboarding/offboarding summary, pendingTasks, formsProgress", "isAuthenticated"),
            ("GET", "/employee/tasks", "Pending employee-role tasks", "—", "Task list", "isAuthenticated"),
            ("POST", "/employee/workflows/:workflowId/tasks/:taskId/complete", "Complete task (simplified)", "{ notes }", "Workflow", "isAuthenticated"),
            ("GET", "/employee/forms/pending", "Pending lifecycle forms", "—", "Form list", "isAuthenticated"),
            ("GET", "/employee/forms/completed", "Completed forms", "—", "Form list", "isAuthenticated"),
            ("GET", "/employee/summary", "Lifecycle summary", "—", "Summary object", "isAuthenticated"),
            ("PUT", "/employee/onboarding/:workflowId/tasks/:taskId", "Employee-only task update", "{ status, notes }", "Workflow", "isAuthenticated + onboarding module; scoped to employeeId"),
            ("PUT", "/employee/onboarding/:workflowId/checklist/:itemId", "Employee checklist", "{ completed }", "Workflow", "Same"),
            ("PUT", "/employee/onboarding/:workflowId/assets/:assetId", "Employee asset acknowledgment", "{ status, notes }", "Workflow", "Same"),
            ("PUT", "/employee/onboarding/:workflowId/training/:trainingId", "Training update", "{ status, progress, notes }", "Workflow", "Same (service uses workflow.training — schema field is trainingModules)"),
            ("POST", "/employee/onboarding/:workflowId/signatures/:signatureId/sign", "Policy sign", "signature payload", "Workflow", "Same"),
            ("PUT", "/employee/onboarding/:workflowId/buddy-checkins/:checkinId", "Buddy check-in", "—", "Workflow", "Same"),
            ("POST", "/employee/onboarding/:workflowId/notes", "Add note", "{ note, visibility? }", "Workflow", "Same"),
            ("GET/POST", "/employee/onboarding/:workflowId/bgv", "BGV via lifecycle path", "—", "Same as both/onboarding BGV", "backgroundVerification feature"),
            ("PATCH", "/employee/onboarding/:workflowId/bgv/checks/:checkId", "Patch BGV check", "—", "Check", "backgroundVerification"),
        ],
    )

    add_heading(doc, "2.3 Related employee APIs (onboarding workflow dependencies)", level=2)
    add_table(
        doc,
        ["Method", "Path", "Purpose", "Auth"],
        [
            ("GET", "/api/v1/both/lifecycle/features", "Org enabled modules/features", "isAuthenticated"),
            ("GET", "/api/v1/both/lifecycle/forms", "Published lifecycle forms", "isAuthenticated"),
            ("GET", "/api/v1/both/lifecycle/forms/:id", "Form definition", "isAuthenticated"),
            ("POST", "/api/v1/both/lifecycle/forms/:id/submit", "Submit form response", "isAuthenticated"),
            ("POST", "/api/v1/both/lifecycle/forms/:id/draft", "Save draft", "isAuthenticated"),
            ("GET", "/api/v1/both/documents/mine", "Employee document library", "isAuthenticated"),
            ("POST", "/api/v1/both/documents/upload", "Upload document (approval flow)", "isAuthenticated + multipart"),
            ("PATCH/DELETE", "/api/v1/both/documents/:id", "Update/delete own document", "isAuthenticated"),
            ("GET/POST/PATCH…", "/api/v1/both/lifecycle/employee/bgv/*", "Full BGV employee portal (upload, drafts, references)", "isAuthenticated + backgroundVerification"),
            ("GET/PUT", "/api/v1/both/profile-details/*", "Profile, job docs, financial/bank details", "isAuthenticatedBoth / employee routes"),
        ],
    )

    add_heading(doc, "2.4 Admin/HR onboarding APIs with no dedicated employee equivalent", level=2)
    add_table(
        doc,
        ["Capability", "Admin/HR path (representative)", "Employee gap"],
        [
            ("Create/update/delete workflow", "POST/PUT/DELETE /admin/onboarding/onboarding/:id", "Employees cannot initiate workflows"),
            ("Templates CRUD / publish / clone", "/admin/onboarding/templates, /admin/lifecycle/templates/onboarding", "No employee template APIs"),
            ("Analytics & org dashboard", "/admin/onboarding/onboarding/analytics", "Employee uses /both/onboarding/dashboard (limited)"),
            ("Workflow stage runtime", "/admin/lifecycle/onboarding/:id/stages/*", "No employee stage read/progress API"),
            ("HR document verify/reject on workflow", "Admin workflow + document hub", "Employee uploads only; verification HR-side"),
            ("Software access create/approve/provision", "POST /admin/lifecycle/onboarding/:id/software-requests", "Employee UI read-only; no employee POST route"),
            ("Milestone / probation HR actions", "/admin/lifecycle/onboarding/:id/milestones/review, probation/*", "No employee milestone submit API in lifecycle employee routes"),
            ("Goals assign / refresh progress", "/admin/lifecycle/onboarding/:id/goals, progress/refresh", "Goals may exist elsewhere (both/goal) but not lifecycle employee assign"),
            ("LMS assign/sync", "/admin/lifecycle/onboarding/:id/lms/*", "Employee completes trainingModules via PUT training only"),
            ("Forms status (HR aggregate)", "GET .../forms/status", "Employee has pending/completed forms lists only"),
            ("Add HR internal notes", "POST /admin/onboarding/onboarding/:id/notes", "Employee can POST note on lifecycle employee route"),
            ("Asset assign (IT)", "PUT /admin/onboarding/onboarding/:id/assets/:assetId", "Employee lifecycle PUT can update status only for own workflow"),
        ],
    )

    # ── 3. Workflow logic ──
    add_heading(doc, "3. Onboarding Steps / Workflow Logic", level=1)
    add_heading(doc, "3.1 Default sequence (backend constants + template seeding)", level=2)
    add_para(
        doc,
        "When createEmployeeOnboardingWorkflow runs (onboardingService.js), artifacts are copied from the org template (or ensureDefaultOnboardingTemplate + onboardingDefaults.js):",
    )
    add_table(
        doc,
        ["Order", "Stage key (template defaults)", "Typical artifacts", "Default mandatory?"],
        [
            ("0", "pre_boarding", "Pre-joining tasks/checklist (feature preBoarding)", "Stage mandatory in DEFAULT_ONBOARDING_STAGES"),
            ("1", "documentation", "Document slots (Aadhaar, PAN, degree, photo, bank, offer ack), employee upload tasks, HR verify task", "Yes"),
            ("2", "it_setup", "Assets (laptop, access card), IT tasks, softwareRequests from packs", "Yes"),
            ("3", "hr_setup", "Payroll task, HR checklist items", "Yes"),
            ("4", "induction", "Training modules, policy signatures, buddy/manager orientation", "Yes"),
        ],
    )
    add_para(
        doc,
        "DEFAULT_ONBOARDING_CHECKLIST items (offer accept, NDA, policies, HR forms, laptop, email, payroll, optional ID card, manager induction, optional buddy) map to featureKeys (offerAcceptance, eSignature, digitalForms, itAssetProvisioning, softwareAccessRequest, payrollSetup, …). "
        "DEFAULT_ONBOARDING_TASKS define employee uploads (identity, bank), HR verify-all-docs, IT laptop, payroll, manager induction, buddy meeting.",
    )

    add_heading(doc, "3.2 Runtime progression rules", level=2)
    add_para(
        doc,
        "Progress: calculateProgressBreakdown weights tasks 22%, documents 18%, checklist 12%, training 8%, policies 8%, assets 8%, forms 5%, goals 5%, bgv 8%, software 6% (feature-gated dimensions zeroed when disabled). workflow.progress is overall 0–100. "
        "Completion: status becomes Completed only when progress === 100, mandatoryFormsCompleted !== false, BGV gate open (Clear/Waived or not blocking), and status in Not Started/In Progress/Overdue.",
    )
    add_para(
        doc,
        "Task dependencies: task.dependencies[] must be Completed before manual completion (getTaskCompletionBlockers). Document-driven tasks auto-sync via synchronizeDocumentDrivenTasks (upload → In Progress; all required docs uploaded → Completed; HR task waits Verified on all mandatory docs). "
        "Template stage dependsOn[] is used in template preview/validation (lifecycleTemplateEngine) but workflow instances do not call initializeOnboardingStages at creation — stage-based unlocking is primarily HR admin runtime, not employee API.",
    )
    add_para(
        doc,
        "No explicit currentStep field on workflow; UI derives from progressBreakdown, open tasks, and checklist. scheduledStartDate + joiningDate control Draft/Scheduled/Not Started via resolveInitialWorkflowStatus.",
    )

    add_heading(doc, "3.3 Optional vs mandatory", level=2)
    add_table(
        doc,
        ["Artifact", "Mandatory rule"],
        [
            ("Template task/checklist/document/asset/training/policy item", "mandatory !== false by default"),
            ("DEFAULT document types", "PASSPORT, EXPERIENCE, RESUME, OTHER optional; Aadhaar, PAN, degree, offer ack, photo, bank required in ONBOARDING_DOCUMENT_TYPES"),
            ("Buddy check-in / ID card checklist", "Optional in defaults"),
            ("Training modules in workflow", "Only instantiated if template entry has videoUrl, documentUrl, or quizQuestions"),
            ("Policy signatures", "Only if template has documentUrl"),
            ("BGV", "Optional feature; may block joining per org bgvPolicy.blockJoiningUntilClear"),
            ("Forms gate", "mandatoryFormsCompleted synced from lifecycleFormGateService"),
        ],
    )

    # ── 4. Frontend ──
    add_heading(doc, "4. Existing Frontend (Employee side)", level=1)
    add_table(
        doc,
        ["Surface", "Route / file", "Status", "APIs used"],
        [
            ("Primary employee UI", "/onboarding → OnboardingRouteWrapper → EmployeeLifecycle → OnboardingHub (employee mode)", "Built", "GET both/onboarding/dashboard, my-workflow; PUT tasks/checklist/training/sign; profile routes for docs"),
            ("Legacy standalone page", "vite_client/src/Employee/Onboarding/MyOnboardingPage.jsx", "Not registered in routes/index.jsx", "GET my-workflow, PUT tasks only; tasks-only UI"),
            ("HR detail reused for employees", "OnboardingHrDetailView sections when employee opens own workflow", "Built (shared)", "Same both/onboarding mutations; deep links to /job-details, /financial-details, /documents"),
            ("BGV panel", "OnboardingBgvPanel.jsx", "Built when feature on", "both/lifecycle/employee/onboarding/:id/bgv + employee/bgv/*"),
            ("Forms panel", "OnboardingFormsPanel.jsx (HR forms/status); employee forms via lifecycle forms API", "Partial for employee", "HR: admin lifecycle forms/status; employee likely forms submit via lifecycle forms routes"),
            ("Access / software", "OnboardingAccessControl.jsx", "Employee: read-only list", "No employee POST; HR uses admin lifecycle software-requests"),
            ("Milestones / goals / LMS / probation panels", "OnboardingMilestonesPanel, Goals, LmsDashboard, ProbationCenter", "HR/admin oriented", "Admin lifecycle endpoints"),
            ("Org tenant wizard", "/organization/onboarding", "Admin only", "organization onboarding APIs — not employee onboarding"),
        ],
    )

    add_heading(doc, "4.1 Design system (as implemented)", level=2)
    add_table(
        doc,
        ["Token", "Expected (PeopleFlow brief)", "Actual in codebase"],
        [
            ("Primary teal", "#1CCA9B", "index.css --color-primary: #20C997; Tailwind theme-primary"),
            ("Coral secondary", "#FF986A", "index.css --color-secondary: #2563eb (blue); --color-accent: #f97316 (orange)"),
            ("Typography", "Manrope", "Manrope only inside .cm-hrms (corporateModern.css) for HR/Admin lifecycle chrome; employee /onboarding shell uses default system fonts + gray-50/page-shell"),
            ("Component library", "—", "Tailwind utility classes; HR views use corporateModern (CmButton, CmRing, …); employee onboarding uses Fa icons, white cards, theme-primary/secondary classes"),
            ("Toast / HTTP", "—", "react-toastify, shared axios instance"),
        ],
    )

    # ── 5. Gaps ──
    add_heading(doc, "5. Gaps & Observations", level=1)
    add_table(
        doc,
        ["Category", "Observation"],
        [
            ("Model without employee API", "workflow.stages runtime (dependsOn, auto-progress) — no schema field on OnboardingWorkflow, initializeOnboardingStages never invoked at create"),
            ("Model without employee API", "milestones[].employeeReview, probation reviews — HR admin endpoints only"),
            ("Model without employee API", "softwareRequests — employee cannot create; only HR admin lifecycle POST"),
            ("Duplicate APIs", "Parallel stacks: /both/onboarding/* vs /both/lifecycle/employee/onboarding/* (different auth middleware and validation strictness)"),
            ("API without UI", "/both/lifecycle/employee/dashboard, /tasks, /summary, /workflows/.../complete — not clearly wired in OnboardingHub (hub uses both/onboarding)"),
            ("API without UI", "MyOnboardingPage.jsx orphaned"),
            ("UI without full API", "Employee Access Control view — no employee provisioning API"),
            ("Naming mismatch", "lifecycleEmployeeService returns training/signatures; schema uses trainingModules/policySignatures (training PUT on lifecycle path likely broken)"),
            ("Naming mismatch", "getEmployeeOnboardingWorkflow populate firstName/lastName; Employee model typically name/lastName"),
            ("Naming mismatch", "Document type enums duplicated (BANK_DETAILS vs Bank Details) in onboardingDocumentSchema"),
            ("Behavior mismatch", "lifecycle employee task update does not call getTaskCompletionBlockers / synchronizeWorkflowDocuments (both/onboarding controller does)"),
            ("Progress UX", "MyOnboardingPage shows all workflow.tasks; both/onboarding sanitizes to role-filtered tasks for employees"),
        ],
    )

    add_heading(doc, "Open Questions", level=1)
    for q in [
        "Should employees use /both/onboarding or /both/lifecycle/employee as the canonical API long term?",
        "Is workflow.stages persistence intentional (missing from Mongoose schema) or is stage runtime HR-only and unfinished for instances?",
        "Which document upload path is authoritative for onboarding: /both/documents/upload, profile job-details/financial-details, or both?",
        "Are milestone employeeReview submissions planned via a public form, or only HR-entered data today?",
        "Should software access requests be employee-initiated (feature softwareAccessRequest) given UI copy implies HR provisions?",
        "Is Organization onboarding wizard completion required before employee onboarding modules activate per tenant?",
        "Confirm brand tokens: migrate employee shell to #1CCA9B / #FF986A / Manrope globally or keep HR-only cm-hrms scope?",
    ]:
        doc.add_paragraph(q, style="List Bullet")

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
